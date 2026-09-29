"""Extract Porsche evidence and assign one editorial owner before B2 source edits."""
import csv, json, re, warnings
from collections import Counter, defaultdict
from pathlib import Path
import openpyxl

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'outputs/b2'
DOWNLOADS=Path('C:/Users/ADMIN/Downloads')
WORKBOOK=DOWNLOADS/'DIGI-TEC_SEO_Keyword_Universe_Dubai.xlsx'
PAGES={p['path']:p for p in json.loads((OUT/'pages-baseline.json').read_text(encoding='utf-8'))}
ROUTES={p['path']:p for p in json.loads((OUT/'route-baseline.json').read_text(encoding='utf-8'))}
HUB='/brands/porsche-service-dubai'
SVC=lambda slug:f'{HUB}/{slug}'

def norm(raw):
    x=re.sub(r'\s+',' ',str(raw or '').casefold().replace('centre','center').replace('servicing','service').replace('gear box','gearbox').replace('diagnostics','diagnostic')).strip()
    return x

def owner_for(raw):
    k=norm(raw)
    # Distinct information, symptom and booking tasks take precedence over nouns.
    if re.search(r'\b(dealer|independent specialist)\b.*\b(vs|versus|comparison)\b|\b(vs|versus)\b.*\b(dealer|independent specialist)\b',k):
        return '/porsche/guides/dealer-vs-independent-specialist','Dealer vs independent specialist','Comparative informational','Existing decision guide; neither owner is a second broad booking page'
    if re.search(r'\b(best|vs|versus|compare|comparison)\b.*(workshop|garage|specialist|dealer)',k):
        return '/best-porsche-workshop-dubai','Workshop comparison','Comparative informational','Existing comparison URL; broad service/repair booking remains at the hub'
    if re.search(r'\b(retrofit|key programming|keyless entry|immobilizer|android auto upgrade|apple carplay upgrade|sound system upgrade|soft close|seat motor|seat repair|boot repair|sunroof|tailgate|window tint)\b',k):
        return '', 'Unsupported equipment/function inquiry','Unverified service','No verified Porsche-specific retrofit or programming scope; do not promise it'
    if re.search(r'\b(window tint|wrap|detailing|ceramic coating|paint protection|ppf|polishing|paint correction)\b',k):
        target='/services/paint-protection-dubai' if 'paint protection' in k or 'ppf' in k else '/services/ceramic-coating' if 'ceramic' in k else '/services/car-polishing-dubai' if 'paint correction' in k or 'polishing' in k or 'detailing' in k else ''
        if target not in PAGES:target=''
        return target,'Paint-care inquiry','Adjacent service' if target else 'Unverified service','Use a verified generic paint-care owner only when the exact service is described'
    if re.search(r'\b(tuning|remap|performance upgrade|ecu tune|stage 1|stage 2)\b',k):
        return '/tuning','Performance project','Cross-brand commercial','Existing tuning service assesses Porsche projects; no new Porsche tuning URL'
    if re.search(r'\b(cost|price|pricing|how much)\b',k):
        return '/porsche/guides/maintenance-cost-dubai','Ownership cost factors','Informational','Explain scope and price drivers without inventing a fixed price'
    if re.search(r'\b(interval|schedule|how often|maintenance plan)\b',k):
        if 'pdk' in k:return '/porsche/guides/pdk-service-intervals','PDK interval planning','Informational','Vehicle-specific maintenance guidance'
        if 'oil' in k:return '/porsche/guides/oil-change-intervals','Oil interval planning','Informational','Vehicle-specific maintenance guidance'
        return '/porsche/guides/service-intervals-uae','Service interval planning','Informational','Vehicle-specific maintenance guidance'
    if re.search(r'\b(what is|explained|how does|how works|meaning|how pdk works)\b',k):
        for term,slug in [('pdk','pdk'),('pasm','pasm'),('pdcc','pdcc'),('pccb','pccb'),('ptm','ptm-awd'),('tiptronic','tiptronic'),('sport chrono','sport-chrono'),('air suspension','air-suspension'),('rear axle steering','rear-axle-steering')]:
            if term in k:return f'/porsche/systems/{slug}',f'{term.upper()} explanation','System informational','Existing system explainer owns operation and fitment'
    problem_rules=[
      (r'cayenne.*air suspension.*(fault|warning|malfunction|problem)', 'cayenne-air-suspension'),
      (r'pasm.*(fault|warning|malfunction)', 'pasm-fault'),
      (r'pdcc.*(fault|warning|malfunction)', 'pdcc-fault'),
      (r'pdk.*(jerk|judder|shudder)', 'pdk-jerking'),
      (r'pdk.*(slip|flare)', 'pdk-slipping'),
      (r'pdk.*(warning|fault|error)', 'pdk-warning-message'),
      (r'(delayed gear|delayed engagement)', 'delayed-gear-engagement'),
      (r'(air suspension warning|suspension warning)', 'air-suspension-warning'),
      (r'(suspension drop|ride height drop)', 'suspension-dropping-overnight'),
      (r'(battery warning|charging warning)', 'battery-warning'),
      (r'(no start|not start|won.t start|starting problem)', 'wont-start'),
      (r'(check engine|engine warning|engine light)', 'check-engine-light'),
      (r'(engine misfire|rough idle|engine shaking)', 'engine-misfire'),
      (r'(engine overheat|overheating)', 'engine-overheating'),
      (r'(coolant leak)', 'coolant-leak'),
      (r'(oil leak)', 'oil-leak'),
      (r'(ac not cool|air conditioning not cool)', 'ac-not-cooling'),
      (r'(brake warning|abs warning)', 'brake-warning-light'),
      (r'(steering vibration|brake vibration)', 'steering-vibration'),
      (r'(pcm.*(not work|black|fault|screen)|screen black|reverse camera not work)', 'check-engine-light'),
    ]
    for pattern,slug in problem_rules:
        if re.search(pattern,k):
            if slug=='check-engine-light' and ('pcm' in k or 'screen' in k or 'camera' in k):
                return SVC('electrical-repair'),'PCM/screen/camera fault','Electrical assessment','No dedicated verified PCM repair owner; assess electrical/infotainment fault before claiming replacement'
            return f'/porsche/problems/{slug}',f'{slug.replace("-"," ")} symptom','Symptom informational','Existing symptom guide owns meaning and triage; service page owns repair booking'
    for term,slug in [('pdk','pdk'),('pasm','pasm'),('pdcc','pdcc'),('pccb','pccb'),('ptm','ptm-awd'),('tiptronic','tiptronic'),('sport chrono','sport-chrono')]:
        if term in k and not re.search(r'\b(repair|replace|service|diagnos|scan)\b',k):
            return f'/porsche/systems/{slug}',f'{term.upper()} explanation','System informational','Existing system guide; repair/warning tasks route separately'
    if re.search(r'\b(992|991|997)\b',k):
        match=re.search(r'\b(992|991|997)\b',k)
        return f'/porsche/911/{match.group(1)}',f'911 {match.group(1)} model service','Model commercial','Generation-specific existing page'
    for pattern,target,label in [
      (r'\b(cayenne)\b','/blog/porsche-cayenne-service-dubai-guide','Cayenne model service'),
      (r'\b(macan)\b','/porsche/macan','Macan model service'),
      (r'\b(panamera)\b','/blog/porsche-panamera-service-dubai-guide','Panamera model service'),
      (r'\b(taycan)\b','/porsche/taycan','Taycan model service'),
      (r'\b(718|boxster|cayman)\b','/porsche/718','718/Boxster/Cayman model service'),
      (r'\b(911)\b','/blog/porsche-911-service-dubai-guide','911 model service')]:
        if re.search(pattern,k) and re.search(r'\b(service|repair|workshop|garage|specialist|maintenance)\b',k):
            return target,label,'Model commercial','Model page serves model ownership; specific repair tasks link to service'
    service_rules=[
      (r'\b(piwis|diagnos\w*|fault code|fault diagnosis|computer scan|control unit scan|dashboard warning)\b','engine-diagnostics','Porsche diagnostics'),
      (r'\b(transmission|gearbox|pdk|clutch|gear shifting)\b','transmission-repair','Porsche transmission/PDK repair'),
      (r'\b(suspension|pasm|pdcc|damper|air strut|bushing|control arm|ride height|shock absorber|strut)\b','suspension-repair','Porsche suspension repair'),
      (r'\b(brake|abs|pccb)\b','brake-repair','Porsche brake repair'),
      (r'\b(oil change|oil service|engine oil|oil filter)\b','oil-change','Porsche oil change'),
      (r'\b(ac |air conditioning|compressor|condenser|refrigerant|cabin filter)\b','ac-repair','Porsche AC repair'),
      (r'\b(battery replac\w*|battery fitting|battery registration)\b','battery-replacement','Porsche battery fitting'),
      (r'\b(electrical|alternator|coding|programming|module|pcm|screen|camera|battery drain|ecu|sensor|starter|wiring|head unit|infotainment|instrument cluster|navigation|stereo|door lock|power window|window regulator)\b','electrical-repair','Porsche electrical assessment'),
      (r'\b(steering|power steering)\b','steering-repair','Porsche steering repair'),
      (r'\b(tire|tyre|wheel|rim|alignment|puncture)\b','tire-repair','Porsche tyre/wheel service'),
      (r'\b(body|collision|accident|paint repair|painting|dent|panel|scratch|bumper)\b','body-repair','Porsche body repair'),
      (r'\b(exhaust|catalytic|emissions?|muffler|oxygen sensor)\b','exhaust-repair','Porsche exhaust repair'),
      (r'\b(fuel|injector|fuel pump)\b','fuel-system-repair','Porsche fuel system repair'),
      (r'\b(engine|mechanical|cooling|coolant|radiator|thermostat|water pump|turbo|misfire|timing belt|timing chain|valve cover|loss of power|spark plug|air filter)\b','mechanical-repair','Porsche mechanical repair'),
    ]
    for pattern,slug,label in service_rules:
        if re.search(pattern,k):
            return SVC(slug),label,'Commercial service','Existing specific service page; exact scope is assessed before work'
    if re.search(r'\b(inspection|pre purchase|ppi)\b',k):
        return '/porsche/guides/pre-purchase-inspection-checklist','Pre-purchase inspection planning','Informational','Checklist and diagnosis scope on existing guide'
    if re.search(r'\b(maintenance)\b',k):
        return '/blog/porsche-maintenance-guide-dubai','Porsche maintenance planning','Informational','Existing planning guide; hub owns booking'
    return HUB,'Porsche service/repair discovery','Broad commercial','Single broad booking and directory owner'

def period(sheet):
    if 'Chart' not in sheet.sheetnames:return ''
    rows=list(sheet['Chart'].values)
    values=[str(r[0])[:10] for r in rows[1:] if r and r[0]]
    return f'{values[0]} to {values[-1]}' if values else ''

records=[]
def add(keyword,source,source_file,source_period,evidence,clicks='',impressions='',ctr='',position='',context=''):
    if not keyword:return
    owner,cluster,intent,reason=owner_for(keyword)
    if owner and owner not in PAGES: raise ValueError(f'Unpublished owner {owner} for {keyword}')
    records.append({'raw_keyword':str(keyword),'normalized_keyword':norm(keyword),'source':source,
      'source_file':source_file,'source_period':source_period,'measured_or_generated':evidence,
      'gsc_clicks':clicks if clicks is not None else '', 'gsc_impressions':impressions if impressions is not None else '',
      'gsc_ctr':ctr if ctr is not None else '', 'gsc_position':position if position is not None else '',
      'search_intent':intent,'cluster':cluster,'primary_owner':owner,
      'secondary_supporting_owner':HUB if owner and owner!=HUB else '',
      'reason':reason,'source_context':context})

book=openpyxl.load_workbook(WORKBOOK,read_only=True,data_only=True)
for rownum,row in enumerate(book['Master Keywords'].values,1):
    if rownum==1 or str(row[1]).casefold()!='porsche':continue
    evidence='MEASURED WORKBOOK' if row[7]=='Existing GSC query' else 'GENERATED TAXONOMY' if row[7]=='Generated taxonomy' else 'RESEARCH KEYWORD'
    add(row[0],f'Master Keywords!A{rownum}',WORKBOOK.name,'2026-04-12 to 2026-09-25' if evidence=='MEASURED WORKBOOK' else '',evidence,*row[8:12],str(row[2]))
for rownum,row in enumerate(book['GSC Opportunities'].values,1):
    if rownum==1 or str(row[1]).casefold()!='porsche':continue
    add(row[0],f'GSC Opportunities!A{rownum}',WORKBOOK.name,'2026-04-12 to 2026-09-25','MEASURED WORKBOOK',*row[2:6],'Opportunity sheet copy')
for rownum,row in enumerate(book['Brand Specific Systems'].values,1):
    if rownum==1 or str(row[0]).casefold()!='porsche':continue
    add('porsche '+str(row[1]),f'Brand Specific Systems!B{rownum}',WORKBOOK.name,'','RESEARCH KEYWORD',context='Brand-specific systems sheet')

# These phrases were supplied as candidate tasks in the B2 brief, not measured
# searches. Keep them separate from Search Console and generated workbook rows.
brief_examples=[
 'what is Porsche PDK','Porsche PDK explained','how PDK works','Porsche PDK jerking',
 'Porsche PDK warning','Porsche transmission warning','Porsche delayed gear engagement',
 'what is Porsche PASM','Porsche PASM explained','Porsche PASM fault','Porsche suspension warning',
 'Porsche PDCC explained','Porsche PCCB explained','Porsche PTM explained',
 'Porsche rear axle steering explained','Porsche Sport Chrono explained',
 'Porsche check engine light','Porsche overheating','Porsche coolant leak','Porsche oil leak',
 'Porsche battery warning','Porsche won\'t start','Porsche AC not cooling',
 'Porsche steering vibration','Porsche brake squeaking','Porsche brake grinding',
 'Porsche loss of power','Porsche turbo malfunction','Porsche PCM not working',
 'Porsche screen black','Porsche reverse camera not working',
 'Porsche 911 service Dubai','Porsche 992 service Dubai','Porsche 991 service Dubai',
 'Porsche 997 service Dubai','Porsche Cayenne service Dubai','Porsche Macan service Dubai',
 'Porsche Panamera service Dubai','Porsche 718 service Dubai','Porsche Taycan service Dubai',
 'Porsche service intervals UAE','Porsche service cost Dubai','Porsche PDK service interval',
 'Porsche oil change interval','Porsche brake replacement cost',
 'Porsche independent specialist vs dealer','Porsche maintenance in UAE heat',
]
for index,term in enumerate(brief_examples,1):
    add(term,f'B2 brief example #{index}','B2 user brief','','OTHER SOURCE — USER BRIEF EXAMPLE',context='Candidate user task, not measured demand')

terms=re.compile(r'\b(porsche|cayenne|macan|panamera|taycan|boxster|cayman|pdk|pasm|pdcc|pccb|piwis)\b',re.I)
gsc_files=[]
for day in ('05','08','13','16','18','19','27'):
    file=DOWNLOADS/f'digitecme.com-Performance-on-Search-2026-09-{day}.xlsx'
    if not file.exists():continue
    gsc_files.append(file)
    w=openpyxl.load_workbook(file,read_only=True,data_only=True)
    span=period(w)
    for rownum,row in enumerate(w['Queries'].values,1):
        if rownum==1 or not row or not terms.search(str(row[0])):continue
        add(row[0],f'Queries!A{rownum}',file.name,span,'MEASURED ORIGINAL GSC',*row[1:5],'Original Search Console query export')

OUT.mkdir(exist_ok=True)
with (OUT/'porsche-keyword-owner-records.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(records[0]));w.writeheader();w.writerows(records)
(OUT/'porsche-keyword-owner-records.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
group=defaultdict(list)
for r in records:group[r['normalized_keyword']].append(r)
conflicts={k:sorted({r['primary_owner'] for r in v}) for k,v in group.items() if len({r['primary_owner'] for r in v})>1}
assert not conflicts,conflicts
owners=Counter(r['primary_owner'] for r in records)
latest=[r for r in records if r['measured_or_generated']=='MEASURED ORIGINAL GSC' and r['source_file'].endswith('2026-09-27.xlsx')]
summary={'source_observations':len(records),'raw_terms':len({r['raw_keyword'] for r in records}),
 'normalized_terms':len(group),'workbook_master_porsche_rows':sum(r['source'].startswith('Master Keywords!') for r in records),
 'original_gsc_records':sum(r['measured_or_generated']=='MEASURED ORIGINAL GSC' for r in records),
 'latest_six_month_queries':len(latest),'latest_six_month_impressions':sum(float(r['gsc_impressions']) for r in latest),
 'latest_six_month_clicks':sum(float(r['gsc_clicks']) for r in latest),
 'evidence_counts':dict(Counter(r['measured_or_generated'] for r in records)),
 'unresolved_owner_conflicts':conflicts,'gsc_files':[p.name for p in gsc_files]}
(OUT/'keyword-summary.json').write_text(json.dumps(summary,indent=2)+'\n',encoding='utf-8')
lines=['# Porsche keyword → primary owner map','',
 'This pre-edit map separates original GSC measurements, measured workbook copies, generated taxonomy, and research. Repeated exports and time windows are separate observations; never add them together. A query owner is an editorial decision, not the URL Google ranked.', '',
 f"{len(records)} source observations; {len(group)} normalized terms; {len(owners)} existing or intentionally unassigned owner classes; zero unresolved primary-owner conflicts.", '',
 '| Primary owner | Source observations | Representative intent |','| --- | ---: | --- |']
for owner,count in owners.most_common():
    example=next(r for r in records if r['primary_owner']==owner)
    lines.append(f"| `{owner or 'NOT TARGETED'}` | {count} | {example['cluster']} |")
lines += ['', 'Every source observation and its rationale is in `porsche-keyword-owner-records.csv`. Empty primary owner means the exact service/function was not verified for Porsche and is deliberately not targeted. Broad workshop booking stays with the hub; `/best-porsche-workshop-dubai` remains a comparison page pending separate consolidation validation.']
(OUT/'porsche-keyword-owner-map.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
print(json.dumps(summary,indent=2))
