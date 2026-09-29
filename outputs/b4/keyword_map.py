"""Separate Ferrari and Lamborghini evidence and assign pre-edit editorial owners."""
import csv, json, re, warnings
from collections import Counter, defaultdict
from pathlib import Path
import openpyxl

warnings.filterwarnings('ignore',category=UserWarning,module='openpyxl')
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'outputs/b4'
DOWNLOADS=Path('C:/Users/ADMIN/Downloads')
WORKBOOK=DOWNLOADS/'DIGI-TEC_SEO_Keyword_Universe_Dubai.xlsx'
PAGES={p['path']:p for p in json.loads((OUT/'pages-baseline.json').read_text(encoding='utf-8'))}
BRANDS=('ferrari','lamborghini')
HUB={b:f'/brands/{b}-service-dubai' for b in BRANDS}
SVC=lambda b,s:f'{HUB[b]}/{s}'

def norm(raw):
    s=str(raw or '').casefold().strip()
    for a,b in [('centre','center'),('servicing','service'),('repairs','repair'),('gear box','gearbox'),('transmission','gearbox')]:s=s.replace(a,b)
    return re.sub(r'\s+',' ',s)

def owner_for(brand,raw):
    k=norm(raw)
    if re.search(r'\b(abu dhabi|sharjah|ajman|al ain)\b',k):
        return '','Outside Dubai service area','Out-of-market','Do not target another city with a Dubai workshop page'
    if re.search(r'\b(best|compare|comparison|versus|vs)\b',k) and re.search(r'workshop|garage|dealer|specialist',k):
        return f'/best-{brand}-workshop-dubai','Workshop selection','Comparison','Selection checklist, distinct from broad service booking'
    if re.search(r'\b(price|pricing|cost|how much|interval|schedule|how often|annual maintenance)\b',k):
        return f'/blog/{brand}-maintenance-guide-dubai','Maintenance planning and cost factors','Informational','Existing guide; no fixed price or universal schedule'
    if re.search(r'\b(key programming|immobilizer|soft close|seat repair|sunroof|tailgate|carplay retrofit|window tint)\b',k):
        return '','Unverified exact function','Unverified service','Keyword evidence does not establish exact workshop capability'
    if re.search(r'\b(tuning|remap|ecu tune|stage 1|stage 2|power upgrade|downpipe)\b',k):
        return '/tuning','Performance project','Cross-brand commercial','Existing tuning enquiry; exact application and gains unpromised'
    if re.search(r'\b(ppf|paint protection|ceramic coating|polishing|paint correction)\b',k):
        target='/services/paint-protection-dubai' if 'ppf' in k or 'paint protection' in k else '/services/ceramic-coating' if 'ceramic' in k else '/services/car-polishing-dubai'
        return target,'Paint/protection task','Cross-brand commercial','Use existing treatment owner; do not create brand × protection URL'
    if brand=='ferrari':
        models=[('purosangue','purosangue'),('portofino','portofino'),('roma','roma'),('sf90','sf90'),('f8','f8-tributo'),('812','812'),('488','488'),('296','296')]
        for term,slug in models:
            if re.search(r'\b'+term+r'\b',k) and re.search(r'\b(service|repair|workshop|garage|maintenance|specialist)\b',k) and not re.search(r'\b(oil|brake|gearbox|suspension|ac|battery|electrical|mechanical|engine|diagnos\w*|body|exhaust|fuel|cooling|screen)\b',k):
                return SVC(brand,slug),f'{term.upper()} model service','Model commercial','Existing distinct Ferrari model owner'
        if re.search(r'\b(458|california|gtc4lusso|f12)\b',k) and re.search(r'\b(service|repair|maintenance)\b',k):
            return HUB[brand],'Other Ferrari model enquiry','Broad hub section','No existing distinct model URL; confirm scope before booking'
    else:
        if re.search(r'\b(urus)\b',k) and re.search(r'\b(service|repair|maintenance|workshop)\b',k) and not re.search(r'\b(oil|brake|gearbox|suspension|ac|battery|electrical|mechanical|engine|diagnos\w*|body|exhaust|fuel|cooling|screen)\b',k):
            return '/blog/lamborghini-urus-service-dubai-guide','Urus model considerations','Model informational','Existing Urus guide supports service pages; no model landing URL'
        if re.search(r'\b(hurac[aá]n|aventador|revuelto|temerario)\b',k) and re.search(r'\b(service|repair|maintenance)\b',k):
            return HUB[brand],'Other Lamborghini model enquiry','Broad hub section','No dedicated model route; inspect variant and supported scope'
    if re.search(r'\b(check engine|gearbox warning|suspension warning|battery warning|won.t start|not start|no start|overheating|coolant leak|oil leak|loss of power|ac not cooling|screen not working|reverse camera not working|jerking|slipping)\b',k):
        if re.search(r'\b(gearbox|jerking|slipping)\b',k):slug='transmission-repair'
        elif re.search(r'\b(suspension)\b',k):slug='suspension-repair'
        elif re.search(r'\b(overheating|coolant leak|oil leak)\b',k):slug='mechanical-repair'
        elif re.search(r'\b(battery|won.t start|not start|no start)\b',k):slug='battery-replacement'
        elif re.search(r'\b(ac not cooling)\b',k):slug='ac-repair'
        elif re.search(r'\b(screen|camera)\b',k):slug='electrical-repair'
        else:slug='engine-diagnostics'
        return SVC(brand,slug),'Symptom assessment','Service section','Existing service supplies diagnostic path; no thin symptom URL'
    services=[
      (r'\b(diagnos\w*|fault code|warning light|computer scan|coding|programming)\b','engine-diagnostics','Diagnostics'),
      (r'\b(gearbox|clutch|shift|f1 gearbox|ldf)\b','transmission-repair','Gearbox repair'),
      (r'\b(suspension|damper|bushing|control arm|ride height|shock absorber|air spring)\b','suspension-repair','Suspension repair'),
      (r'\b(brake|abs|carbon ceramic)\b','brake-repair','Brake repair'),
      (r'\b(oil change|oil service|engine oil|oil filter)\b','oil-change','Oil service'),
      (r'\b(ac|air conditioning|compressor|condenser|refrigerant|climate control)\b','ac-repair','AC repair'),
      (r'\b(battery|charging|starter|no start|battery drain)\b','battery-replacement','12V battery assessment'),
      (r'\b(electrical|alternator|module|sensor|wiring|ecu|screen|camera|infotainment)\b','electrical-repair','Electrical assessment'),
      (r'\b(steering|power steering)\b','steering-repair','Steering repair'),
      (r'\b(tire|tyre|wheel|rim|alignment|puncture)\b','tire-repair','Tyre/wheel service'),
      (r'\b(body|collision|accident|paint repair|painting|dent|panel|bumper)\b','body-repair','Body repair'),
      (r'\b(exhaust|catalytic|emissions|muffler)\b','exhaust-repair','Exhaust repair'),
      (r'\b(fuel|injector|fuel pump)\b','fuel-system-repair','Fuel repair'),
      (r'\b(engine|mechanical|cooling|coolant|radiator|thermostat|water pump|turbo|belt|chain|leak)\b','mechanical-repair','Mechanical repair'),
    ]
    for pattern,slug,label in services:
        if re.search(pattern,k):return SVC(brand,slug),f'{brand.title()} {label}','Commercial service','Existing brand-specific service; scope confirmed for vehicle'
    if re.search(r'\b(maintenance|ownership)\b',k):
        return f'/blog/{brand}-maintenance-guide-dubai','Maintenance planning','Informational','Guide supports broad hub booking'
    return HUB[brand],f'{brand.title()} service/repair discovery','Broad commercial','Single broad booking and directory owner'

def period(w):
    if 'Chart' not in w.sheetnames:return ''
    v=[str(r[0])[:10] for r in list(w['Chart'].values)[1:] if r and r[0]]
    return f'{v[0]} to {v[-1]}' if v else ''

records={b:[] for b in BRANDS}
def add(brand,keyword,source,file,span,evidence,clicks='',impressions='',ctr='',position='',context=''):
    if not keyword:return
    owner,cluster,intent,reason=owner_for(brand,keyword)
    if owner and owner not in PAGES:raise ValueError(f'Unpublished owner {owner}: {keyword}')
    records[brand].append(dict(raw_keyword=str(keyword),normalized_keyword=norm(keyword),brand=brand.title(),source=source,source_file=file,source_period=span,measured_or_generated=evidence,gsc_clicks=clicks if clicks is not None else '',gsc_impressions=impressions if impressions is not None else '',gsc_ctr=ctr if ctr is not None else '',gsc_position=position if position is not None else '',search_intent=intent,cluster=cluster,primary_owner=owner,secondary_supporting_owner=HUB[brand] if owner and owner!=HUB[brand] else '',coverage_type='NOT TARGETED — INTENTIONALLY' if not owner else 'EXISTING OWNER',reason=reason,source_context=context))

w=openpyxl.load_workbook(WORKBOOK,read_only=True,data_only=True)
for n,row in enumerate(w['Master Keywords'].values,1):
    if n==1 or str(row[1]).casefold() not in BRANDS:continue
    b=str(row[1]).casefold()
    evidence='MEASURED WORKBOOK' if row[7]=='Existing GSC query' else 'GENERATED TAXONOMY' if row[7]=='Generated taxonomy' else 'RESEARCH KEYWORD'
    add(b,row[0],f'Master Keywords!A{n}',WORKBOOK.name,'2026-04-12 to 2026-09-25' if evidence=='MEASURED WORKBOOK' else '',evidence,*row[8:12],str(row[2]))
for n,row in enumerate(w['GSC Opportunities'].values,1):
    if n==1 or str(row[1]).casefold() not in BRANDS:continue
    b=str(row[1]).casefold()
    add(b,row[0],f'GSC Opportunities!A{n}',WORKBOOK.name,'2026-04-12 to 2026-09-25','MEASURED WORKBOOK',*row[2:6],'Opportunity sheet copy')
for n,row in enumerate(w['Brand Specific Systems'].values,1):
    if n==1 or str(row[0]).casefold() not in BRANDS:continue
    b=str(row[0]).casefold()
    add(b,f'{b} {row[1]}',f'Brand Specific Systems!B{n}',WORKBOOK.name,'','RESEARCH KEYWORD',context='System research suggestion')

brief={
 'ferrari':['Ferrari service Dubai','Ferrari repair Dubai','Ferrari workshop Dubai','Ferrari diagnostics Dubai','Ferrari engine repair Dubai','Ferrari transmission repair Dubai','Ferrari gearbox warning','Ferrari suspension repair Dubai','Ferrari brake repair Dubai','Ferrari AC repair Dubai','Ferrari battery replacement','Ferrari overheating','Ferrari coolant leak','Ferrari oil leak','Ferrari check engine light','Ferrari screen repair','Ferrari 296 GTB service Dubai','Ferrari SF90 service Dubai','Ferrari Roma service Dubai','Ferrari Purosangue service Dubai','Ferrari 812 service Dubai','Ferrari 488 service Dubai','Ferrari service cost Dubai','Ferrari maintenance interval'],
 'lamborghini':['Lamborghini service Dubai','Lamborghini repair Dubai','Lamborghini workshop Dubai','Lamborghini diagnostics Dubai','Lamborghini engine repair Dubai','Lamborghini transmission repair Dubai','Lamborghini gearbox warning','Lamborghini suspension repair Dubai','Lamborghini brake repair Dubai','Lamborghini AC repair Dubai','Lamborghini battery replacement','Lamborghini overheating','Lamborghini coolant leak','Lamborghini oil leak','Lamborghini check engine light','Lamborghini reverse camera repair','Lamborghini Urus service Dubai','Lamborghini Huracan service Dubai','Lamborghini Aventador service Dubai','Lamborghini Revuelto service Dubai','Lamborghini Temerario service Dubai','Lamborghini service cost Dubai','Lamborghini maintenance interval'],
}
for b,terms in brief.items():
    for n,term in enumerate(terms,1):add(b,term,f'B4 brief example #{n}','B4 user brief','','OTHER SOURCE — USER BRIEF EXAMPLE',context='Candidate user task, not measured demand')

patterns={'ferrari':re.compile(r'\b(ferrari|purosangue|portofino|sf90|f8 tributo)\b',re.I),'lamborghini':re.compile(r'\b(lamborghini|hurac[aá]n|aventador|revuelto|temerario)\b',re.I)}
gsc_files=[]
for day in ('05','08','13','16','18','19','27'):
    file=DOWNLOADS/f'digitecme.com-Performance-on-Search-2026-09-{day}.xlsx'
    if not file.exists():continue
    gsc_files.append(file)
    book=openpyxl.load_workbook(file,read_only=True,data_only=True)
    span=period(book)
    for n,row in enumerate(book['Queries'].values,1):
        if n==1 or not row:continue
        for b,pattern in patterns.items():
            if pattern.search(str(row[0])):add(b,row[0],f'Queries!A{n}',file.name,span,'MEASURED ORIGINAL GSC',*row[1:5],'Original GSC query export')

summary={}
for b,items in records.items():
    with (OUT/f'{b}-keyword-owner-records.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(items[0]));writer.writeheader();writer.writerows(items)
    (OUT/f'{b}-keyword-owner-records.json').write_text(json.dumps(items,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    group=defaultdict(list)
    for r in items:group[r['normalized_keyword']].append(r)
    conflicts={k:sorted({r['primary_owner'] for r in v}) for k,v in group.items() if len({r['primary_owner'] for r in v})>1}
    if conflicts:raise ValueError(f'{b} owner conflicts: {conflicts}')
    owners=Counter(r['primary_owner'] for r in items)
    latest=[r for r in items if r['measured_or_generated']=='MEASURED ORIGINAL GSC' and r['source_file'].endswith('2026-09-27.xlsx')]
    summary[b]={'source_observations':len(items),'raw_terms':len({r['raw_keyword'] for r in items}),'normalized_terms':len(group),'master_rows':sum(r['source'].startswith('Master Keywords!') for r in items),'original_gsc_records':sum(r['measured_or_generated']=='MEASURED ORIGINAL GSC' for r in items),'original_gsc_normalized_terms':len({r['normalized_keyword'] for r in items if r['measured_or_generated']=='MEASURED ORIGINAL GSC'}),'latest_six_month_queries':len(latest),'latest_six_month_impressions':sum(float(r['gsc_impressions']) for r in latest),'latest_six_month_clicks':sum(float(r['gsc_clicks']) for r in latest),'evidence_counts':dict(Counter(r['measured_or_generated'] for r in items)),'unresolved_owner_conflicts':conflicts}
    lines=[f'# {b.title()} keyword → primary owner map','','Pre-edit editorial map. Original GSC, workbook copies, generated taxonomy, research and brief examples stay separate. Overlapping GSC export periods are never summed. A primary owner is not evidence of the URL Google ranked.','',f'{len(items)} observations; {len(group)} normalized terms; zero unresolved owner conflicts.','','| Primary owner | Observations | Representative task |','| --- | ---: | --- |']
    for owner,count in owners.most_common():
        example=next(r for r in items if r['primary_owner']==owner)
        lines.append(f"| `{owner or 'NOT TARGETED'}` | {count} | {example['cluster']} |")
    lines.extend(['',f'Each observation and reason is in `{b}-keyword-owner-records.csv`. Blank owner means an unverified exact service or out-of-market query. The hub owns broad booking; the best-workshop URL keeps a selection task pending future evidence-based consolidation.'])
    (OUT/f'{b}-keyword-owner-map.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
(OUT/'keyword-summary.json').write_text(json.dumps(summary,indent=2)+'\n',encoding='utf-8')
print(json.dumps(summary,indent=2))
