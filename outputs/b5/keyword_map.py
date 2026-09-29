"""B5 provenance-preserving keyword extraction and separate editorial owner maps."""
import csv,json,re,warnings
from collections import Counter,defaultdict
from pathlib import Path
import openpyxl
warnings.filterwarnings('ignore',category=UserWarning,module='openpyxl')
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'outputs/b5';DOWNLOADS=Path('C:/Users/ADMIN/Downloads')
WORKBOOK=DOWNLOADS/'DIGI-TEC_SEO_Keyword_Universe_Dubai.xlsx'
PAGES={p['path']:p for p in json.loads((OUT/'pages-baseline.json').read_text(encoding='utf-8'))}
BRANDS=('rolls-royce','bentley','maybach')
HUB={b:f'/brands/{b}-service-dubai' for b in BRANDS}
GUIDE={'rolls-royce':'/blog/rolls-royce-best-workshop-dubai','bentley':'/blog/bentley-best-workshop-dubai','maybach':'/blog/maybach-best-workshop-dubai'}
MODEL={'rolls-royce':'/blog/rolls-royce-ghost-service-dubai-guide','bentley':'/blog/bentley-continental-gt-service-dubai-guide','maybach':'/blog/maybach-s580-service-dubai-guide'}
SVC=lambda b,s:f'{HUB[b]}/{s}'
def norm(raw):
    s=str(raw or '').casefold().strip().replace('rolls royce','rolls-royce').replace('mercedes maybach','mercedes-maybach')
    s=s.replace('maybach mercedes','mercedes-maybach')
    for a,b in [('centre','center'),('servicing','service'),('repairs','repair'),('gear box','gearbox'),('transmission','gearbox')]:s=s.replace(a,b)
    return re.sub(r'\s+',' ',s)
def owner_for(brand,raw):
    k=norm(raw)
    if re.search(r'\b(abu dhabi|sharjah|ajman|al ain)\b',k):return '','Outside Dubai service area','Out-of-market','Do not target another city with a Dubai workshop page'
    if brand=='rolls-royce' and 'spectre' in k and re.search(r'\b(engine oil|oil change|spark plug|turbo|engine repair)\b',k):return '','Inapplicable Spectre ICE service','Incompatible task','Spectre is electric; do not assign combustion-engine service'
    if re.search(r'\b(key programming|immobilizer|soft close|seat repair|sunroof|tailgate)\b',k):return '','Unverified exact function','Unverified service','Keyword evidence does not establish exact workshop capability'
    if re.search(r'\b(best|compare|comparison|versus|vs)\b',k) and re.search(r'workshop|garage|dealer|specialist',k):return GUIDE[brand],'Workshop selection','Comparison','Existing selection guide; distinct from broad booking'
    if re.search(r'\b(price|pricing|cost|how much|interval|schedule|how often|annual maintenance)\b',k):return GUIDE[brand],'Maintenance/cost planning','Informational','Existing informational guide explains cost factors and vehicle-specific schedules'
    if re.search(r'\b(tuning|remap|ecu tune|stage 1|stage 2|power upgrade|downpipe)\b',k):return '/tuning','Performance project','Cross-brand commercial','Existing tuning enquiry; exact application and gains unpromised'
    if re.search(r'\b(ppf|paint protection|ceramic coating|polishing|paint correction)\b',k):
        target='/services/paint-protection-dubai' if 'ppf' in k or 'paint protection' in k else '/services/ceramic-coating' if 'ceramic' in k else '/services/car-polishing-dubai'
        return target,'Paint/protection task','Cross-brand commercial','Existing treatment owner; no brand × protection URL'
    if brand=='rolls-royce':
        if 'ghost' in k and re.search(r'\b(service|repair|workshop|garage|maintenance)\b',k) and not re.search(r'\b(oil|brake|gearbox|suspension|ac|battery|electrical|engine|diagnos\w*)\b',k):return MODEL[brand],'Ghost model service','Model informational','Existing Ghost guide and contextual service links'
        if re.search(r'\b(cullinan|phantom|spectre|wraith|dawn)\b',k) and re.search(r'\b(service|repair|maintenance|workshop)\b',k):return HUB[brand],'Other Rolls-Royce model enquiry','Broad hub section','No dedicated model route; confirm exact powertrain and service scope'
    if brand=='bentley':
        if re.search(r'\b(continental gt|continental gtc)\b',k) and re.search(r'\b(service|repair|workshop|garage|maintenance)\b',k) and not re.search(r'\b(oil|brake|gearbox|suspension|ac|battery|electrical|engine|diagnos\w*)\b',k):return MODEL[brand],'Continental GT model service','Model informational','Existing Continental GT guide'
        if re.search(r'\b(bentayga|flying spur|mulsanne)\b',k) and re.search(r'\b(service|repair|maintenance|workshop)\b',k):return HUB[brand],'Other Bentley model enquiry','Broad hub section','No dedicated model route; service scope confirmed per vehicle'
    if brand=='maybach':
        if re.search(r'\b(s580|s680|s.class)\b',k) and re.search(r'\b(service|repair|workshop|garage|maintenance)\b',k) and not re.search(r'\b(oil|brake|gearbox|suspension|ac|battery|electrical|engine|diagnos\w*)\b',k):return MODEL[brand],'Mercedes-Maybach S-Class service','Model informational','Existing S580 guide supports Maybach commercial owners, not generic Mercedes intent'
        if re.search(r'\b(gls|gls 600)\b',k) and re.search(r'\b(service|repair|maintenance|workshop)\b',k):return HUB[brand],'Mercedes-Maybach GLS enquiry','Broad hub section','No separate Maybach GLS route; distinguish from generic Mercedes GLS'
    if brand=='bentley' and re.search(r'\b(reverse camera|reversing camera|backup camera|camera fault|camera not working)\b',k):return SVC(brand,'electrical-repair'),'Bentley camera assessment','Electrical service section','Existing reverse-camera section owns fault and installation enquiries; verify requested scope'
    if re.search(r'\b(check engine|gearbox warning|suspension warning|battery warning|won.t start|not start|no start|overheating|coolant leak|oil leak|loss of power|ac not cooling|screen not working|camera not working|jerking|slipping|dropping overnight|car dropping)\b',k):
        if re.search(r'\b(gearbox|jerking|slipping)\b',k):slug='transmission-repair'
        elif re.search(r'\b(suspension|dropping)\b',k):slug='suspension-repair'
        elif re.search(r'\b(overheating|coolant leak|oil leak)\b',k):slug='mechanical-repair'
        elif re.search(r'\b(battery|won.t start|not start|no start)\b',k):slug='battery-replacement'
        elif 'ac not cooling' in k:slug='ac-repair'
        elif re.search(r'\b(screen|camera)\b',k):slug='electrical-repair'
        else:slug='engine-diagnostics'
        return SVC(brand,slug),'Symptom assessment','Service section','Existing service provides diagnostic path; no thin symptom URL'
    services=[
      (r'\b(diagnos\w*|fault code|warning light|computer scan|xentry|coding|programming)\b','engine-diagnostics','Diagnostics'),
      (r'\b(gearbox|clutch|shift|9g.tronic)\b','transmission-repair','Gearbox repair'),
      (r'\b(suspension|airmatic|damper|bushing|control arm|ride height|shock absorber|air spring|compressor)\b','suspension-repair','Suspension repair'),
      (r'\b(brake|abs|carbon ceramic)\b','brake-repair','Brake repair'),
      (r'\b(oil change|oil service|engine oil|oil filter)\b','oil-change','Oil service'),
      (r'\b(ac|air conditioning|condenser|refrigerant|climate control)\b','ac-repair','AC repair'),
      (r'\b(battery|charging|starter|no start|battery drain|12v)\b','battery-replacement','Low-voltage battery assessment'),
      (r'\b(electrical|alternator|module|sensor|wiring|ecu|screen|camera|infotainment|mbux|comand|audio)\b','electrical-repair','Electrical assessment'),
      (r'\b(steering|power steering)\b','steering-repair','Steering repair'),
      (r'\b(tire|tyre|wheel|rim|alignment|puncture)\b','tire-repair','Tyre/wheel service'),
      (r'\b(body|collision|accident|paint repair|painting|dent|panel|bumper)\b','body-repair','Body repair'),
      (r'\b(exhaust|catalytic|emissions|muffler)\b','exhaust-repair','Exhaust repair'),
      (r'\b(fuel|injector|fuel pump)\b','fuel-system-repair','Fuel repair'),
      (r'\b(engine|mechanical|cooling|coolant|radiator|thermostat|water pump|turbo|belt|chain|leak)\b','mechanical-repair','Mechanical repair'),
    ]
    for pattern,slug,label in services:
        if re.search(pattern,k):return SVC(brand,slug),f'{brand.title()} {label}','Commercial service','Existing brand-specific service; exact vehicle scope confirmed'
    if re.search(r'\b(maintenance|ownership)\b',k):return GUIDE[brand],'Maintenance planning','Informational','Existing guide supports broad hub booking'
    return HUB[brand],f'{brand.title()} service/repair discovery','Broad commercial','One broad booking and directory owner'
def period(w):
    if 'Chart' not in w.sheetnames:return ''
    v=[str(r[0])[:10] for r in list(w['Chart'].values)[1:] if r and r[0]]
    return f'{v[0]} to {v[-1]}' if v else ''
records={b:[] for b in BRANDS}
def add(brand,keyword,source,file,span,evidence,clicks='',impressions='',ctr='',position='',context=''):
    if not keyword:return
    owner,cluster,intent,reason=owner_for(brand,keyword)
    if owner and owner not in PAGES:raise ValueError(f'Unpublished owner {owner}: {keyword}')
    records[brand].append(dict(raw_keyword=str(keyword),normalized_keyword=norm(keyword),brand={'rolls-royce':'Rolls-Royce','bentley':'Bentley','maybach':'Maybach'}[brand],source=source,source_file=file,source_period=span,measured_or_generated=evidence,gsc_clicks=clicks if clicks is not None else '',gsc_impressions=impressions if impressions is not None else '',gsc_ctr=ctr if ctr is not None else '',gsc_position=position if position is not None else '',search_intent=intent,cluster=cluster,primary_owner=owner,secondary_supporting_owner=HUB[brand] if owner and owner!=HUB[brand] else '',coverage_type='NOT TARGETED — INTENTIONALLY' if not owner else 'EXISTING OWNER',reason=reason,source_context=context))
def brand_id(label):return {'Rolls-Royce':'rolls-royce','Bentley':'bentley','Mercedes-Maybach':'maybach'}.get(str(label))
w=openpyxl.load_workbook(WORKBOOK,read_only=True,data_only=True)
for n,row in enumerate(w['Master Keywords'].values,1):
    if n==1 or not brand_id(row[1]):continue
    b=brand_id(row[1]);e='MEASURED WORKBOOK' if row[7]=='Existing GSC query' else 'GENERATED TAXONOMY' if row[7]=='Generated taxonomy' else 'RESEARCH KEYWORD'
    add(b,row[0],f'Master Keywords!A{n}',WORKBOOK.name,'2026-04-12 to 2026-09-25' if e=='MEASURED WORKBOOK' else '',e,*row[8:12],str(row[2]))
for n,row in enumerate(w['GSC Opportunities'].values,1):
    if n==1 or not brand_id(row[1]):continue
    b=brand_id(row[1]);add(b,row[0],f'GSC Opportunities!A{n}',WORKBOOK.name,'2026-04-12 to 2026-09-25','MEASURED WORKBOOK',*row[2:6],'Opportunity sheet copy')
for n,row in enumerate(w['Brand Specific Systems'].values,1):
    if n==1 or not brand_id(row[0]):continue
    b=brand_id(row[0]);add(b,f'{b} {row[1]}',f'Brand Specific Systems!B{n}',WORKBOOK.name,'','RESEARCH KEYWORD',context='System research suggestion')
brief={
 'rolls-royce':['Rolls Royce service Dubai','Rolls-Royce repair Dubai','Rolls Royce workshop Dubai','Rolls Royce diagnostics Dubai','Rolls Royce suspension repair Dubai','Rolls Royce air suspension repair','Rolls Royce transmission repair Dubai','Rolls Royce brake repair Dubai','Rolls Royce AC repair Dubai','Rolls Royce battery replacement','Rolls Royce camera fault','Rolls Royce Ghost service Dubai','Rolls Royce Cullinan service Dubai','Rolls Royce Spectre service Dubai','Rolls Royce service cost Dubai','Rolls Royce maintenance interval'],
 'bentley':['Bentley service Dubai','Bentley repair Dubai','Bentley workshop Dubai','Bentley diagnostics Dubai','Bentley reverse camera Dubai','Bentley reverse camera not working','Bentley camera repair Dubai','Bentley transmission repair Dubai','Bentley suspension repair Dubai','Bentley brake repair Dubai','Bentley AC repair Dubai','Bentley battery replacement','Bentley Continental GT service Dubai','Bentley Bentayga service Dubai','Bentley service cost Dubai','Bentley maintenance interval'],
 'maybach':['Maybach service Dubai','Mercedes Maybach repair Dubai','Maybach workshop Dubai','Maybach diagnostics Dubai','Maybach AIRMATIC repair Dubai','Maybach suspension warning','Maybach transmission repair Dubai','Maybach MBUX repair','Maybach screen repair','Maybach brake repair Dubai','Maybach AC repair Dubai','Maybach battery replacement','Maybach S580 service Dubai','Maybach S680 service Dubai','Maybach GLS 600 service Dubai','Maybach service cost Dubai'],
}
for b,terms in brief.items():
    for n,term in enumerate(terms,1):add(b,term,f'B5 brief example #{n}','B5 user brief','','OTHER SOURCE — USER BRIEF EXAMPLE',context='Candidate task, not measured demand')
patterns={'rolls-royce':re.compile(r'\brolls[ -]royce\b',re.I),'bentley':re.compile(r'\bbentley\b',re.I),'maybach':re.compile(r'\bmaybach\b',re.I)}
for day in ('05','08','13','16','18','19','27'):
    file=DOWNLOADS/f'digitecme.com-Performance-on-Search-2026-09-{day}.xlsx'
    if not file.exists():continue
    book=openpyxl.load_workbook(file,read_only=True,data_only=True);span=period(book)
    for n,row in enumerate(book['Queries'].values,1):
        if n==1 or not row:continue
        for b,pattern in patterns.items():
            if pattern.search(str(row[0])):add(b,row[0],f'Queries!A{n}',file.name,span,'MEASURED ORIGINAL GSC',*row[1:5],'Original GSC query export')
summary={}
for b,items in records.items():
    with (OUT/f'{b}-keyword-owner-records.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(items[0]));writer.writeheader();writer.writerows(items)
    (OUT/f'{b}-keyword-owner-records.json').write_text(json.dumps(items,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    group=defaultdict(list)
    for r in items:group[r['normalized_keyword']].append(r)
    conflicts={k:sorted({r['primary_owner'] for r in v}) for k,v in group.items() if len({r['primary_owner'] for r in v})>1}
    if conflicts:raise ValueError(f'{b} owner conflicts: {conflicts}')
    owners=Counter(r['primary_owner'] for r in items)
    latest=[r for r in items if r['measured_or_generated']=='MEASURED ORIGINAL GSC' and r['source_file'].endswith('2026-09-27.xlsx')]
    summary[b]={'source_observations':len(items),'raw_terms':len({r['raw_keyword'] for r in items}),'normalized_terms':len(group),'master_rows':sum(r['source'].startswith('Master Keywords!') for r in items),'original_gsc_records':sum(r['measured_or_generated']=='MEASURED ORIGINAL GSC' for r in items),'original_gsc_normalized_terms':len({r['normalized_keyword'] for r in items if r['measured_or_generated']=='MEASURED ORIGINAL GSC'}),'latest_six_month_queries':len(latest),'latest_six_month_impressions':sum(float(r['gsc_impressions']) for r in latest),'latest_six_month_clicks':sum(float(r['gsc_clicks']) for r in latest),'evidence_counts':dict(Counter(r['measured_or_generated'] for r in items)),'unresolved_owner_conflicts':conflicts}
    lines=[f'# {b.title()} keyword → primary owner map','','Pre-edit editorial map. Original GSC, workbook copies, generated taxonomy, research and brief examples stay separate. Overlapping GSC export periods are never summed. A primary owner is not evidence of the URL Google ranked.','',f'{len(items)} observations; {len(group)} normalized terms; zero unresolved owner conflicts.','','| Primary owner | Observations | Representative task |','| --- | ---: | --- |']
    for owner,count in owners.most_common():
        example=next(r for r in items if r['primary_owner']==owner)
        lines.append(f"| `{owner or 'NOT TARGETED'}` | {count} | {example['cluster']} |")
    lines.extend(['',f'Each observation and reason is in `{b}-keyword-owner-records.csv`. Blank owner means an unverified exact service, incompatible task or out-of-market query. Maybach rows require explicit Maybach evidence; generic Mercedes queries are excluded.'])
    (OUT/f'{b}-keyword-owner-map.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
(OUT/'keyword-summary.json').write_text(json.dumps(summary,indent=2)+'\n',encoding='utf-8')
print(json.dumps(summary,indent=2))
