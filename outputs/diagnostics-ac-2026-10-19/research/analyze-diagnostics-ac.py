"""Read-only GSC extraction; no source workbook is saved."""
from pathlib import Path
import csv, hashlib, json, re, warnings, sys
import openpyxl

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]
SOURCES = Path(r'C:\Users\ADMIN\Downloads')
warnings.filterwarnings('ignore', message='Workbook contains no default style')
sys.stdout.reconfigure(encoding='utf-8')

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(q): return re.sub(r'\s+', ' ', re.sub(r'[-–—]', ' ', q.lower())).strip()
def write_json(name, data): (OUT/name).write_text(json.dumps(data,indent=2,ensure_ascii=False),encoding='utf-8')
def write_csv(name, rows):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]) if rows else []); w.writeheader(); w.writerows(rows)
def readable(s):
    replacements={'the92':'the 92','current92':'current 92','primary92':'primary 92','Primary92':'Primary 92','Historical28':'Historical 28','Last24hours':'Last 24 hours','export19October':'export 19 October','18October06':'18 October 06','19October05':'19 October 05','October2026':'October 2026','September2026':'September 2026','-17original':'- 17 original','Historical7':'Historical 7','older92':'older 92','167-day':'167-day','Some1000':'Some 1000','exported(as':'exported (as','position(especially6–20)':'position (especially 6–20)','generalAC':'general AC','RangeRoverAC':'Range Rover AC','PorscheAC':'Porsche AC','Top20':'Top 20','the92-day':'the 92-day','aDIP2':'a DIP2','variants(XENTRY':'variants (XENTRY','nonoverlapping28':'nonoverlapping 28','query-level28':'query-level 28','position6.60':'average position 6.60','on63impressions':'on 63 impressions','65impressions':'65 impressions','49impressions':'49 impressions','42impressions,1click and19.86':'42 impressions, 1 click and 19.86','38impressions':'38 impressions','28impressions':'28 impressions','62impressions':'62 impressions','90impressions':'90 impressions','111impressions':'111 impressions','26impressions':'26 impressions','220impressions':'220 impressions','61impressions':'61 impressions','82impressions':'82 impressions','74impressions':'74 impressions','64impressions':'64 impressions','35impressions':'35 impressions','71impressions':'71 impressions','19impressions':'19 impressions','at15.':'at 15.','at13.':'at 13.','at16.':'at 16.','at20.':'at 20.','at25.':'at 25.','at28.':'at 28.','at18.':'at 18.','at24.':'at 24.','at35.':'at 35.','at38.':'at 38.','at36.':'at 36.','at34.':'at 34.','at37.':'at 37.','at19.':'at 19.','at27.':'at 27.','at14.':'at 14.','plus5fresh24himpressions':'plus 5 fresh 24-hour impressions'}
    for a,b in replacements.items(): s=s.replace(a,b)
    return s
def related(q):
    return bool(re.search(r'\bdiagnos|\bac\b|air\s*con|aircon|air conditioning|\b(?:xentry|ista|piwis|odis|epc|misfire)\b|check engine|warning light|no start|won.t start|warm air|refrigerant|compressor|condenser|evaporator|gas leak|not cooling|drivetrain malfunction|airmatic malfunction|fault (?:code|diagnos)|تشخيص|مكيف|فحص كمبيوتر',norm(q)))

books={}; manifest=[]; evidence=[]; pages=[]
seen_data={}
for p in sorted(SOURCES.glob('digitecme.com-Performance-on-Search-2026-*.xlsx')):
    before=sha(p); wb=openpyxl.load_workbook(p,read_only=True,data_only=True); data={}; ranges={}
    for s in wb:
        v=list(s.values); headers=list(v[0]) if v else []
        data[s.title]=[dict(zip(headers,row)) for row in v[1:] if any(x is not None for x in row)]
        ranges[s.title]=f'A1:{openpyxl.utils.get_column_letter(s.max_column)}{s.max_row}'
    wb.close(); assert before==sha(p)
    filt={r.get('Filter'):r.get('Value') for r in data.get('Filters',[])}
    tk=next((k for k in (data.get('Chart') or [{}])[0] if k=='Date' or k.startswith('Time (')),None)
    dates=[str(r[tk]) for r in data.get('Chart',[]) if tk and r.get(tk)]
    source={'path':str(p),'sha256':before,'source_unchanged':True,'filters':filt,'time_column':tk,'start':min(dates) if dates else None,'end':max(dates) if dates else None,'time_rows':len(dates),'sheets':{s:{'range':ranges[s],'rows':len(rows)} for s,rows in data.items()},'data_table_sha256':hashlib.sha256(json.dumps(data,sort_keys=True,ensure_ascii=False).encode()).hexdigest()}
    source['semantic_duplicate_of']=seen_data.get(source['data_table_sha256'])
    seen_data.setdefault(source['data_table_sha256'],p.name)
    books[p.name]={'source':source,'data':data}; manifest.append(source)
    common={'source_file':p.name,'source_path':str(p),'source_period':f"{source['start']} to {source['end']}",'filters':json.dumps(filt,ensure_ascii=False,separators=(',',':'))}
    for i,r in enumerate(data.get('Queries',[]),2):
        if related(r['Top queries']): evidence.append({**common,'query':r['Top queries'],'clicks':r['Clicks'],'impressions':r['Impressions'],'ctr':r['CTR'],'position':r['Position'],'source_range':f'Queries!A{i}:E{i}'})
    for i,r in enumerate(data.get('Pages',[]),2):
        if related(r['Top pages']): pages.append({**common,'url':r['Top pages'],'clicks':r['Clicks'],'impressions':r['Impressions'],'ctr':r['CTR'],'position':r['Position'],'source_range':f'Pages!A{i}:E{i}','query_attribution':'Unavailable: independent Pages tab, not query-by-page evidence'})
write_json('source-manifest.json',manifest); write_json('extracted-workbooks.json',books)
write_csv('diagnostics-ac-query-evidence-all-periods.csv',evidence);write_csv('diagnostics-ac-independent-page-evidence.csv',pages)
PRIMARY={'24h':'digitecme.com-Performance-on-Search-2026-10-19.xlsx','28d_historical':'digitecme.com-Performance-on-Search-2026-09-05.xlsx','92d':'digitecme.com-Performance-on-Search-2026-10-12 (1).xlsx'}
lookups={label:{norm(r['Top queries']):(i,r) for i,r in enumerate(books[name]['data']['Queries'],2)} for label,name in PRIMARY.items()}
aud=json.loads((ROOT/'outputs/diagnostics-ac-2026-10-19/audit/baseline-owner-audit.json').read_text(encoding='utf-8'))
owners={r['route']:r for r in aud['records']}
focus={'/services/car-diagnostics-dubai','/services/car-ac-repair-dubai','/brands/bmw-service-dubai/engine-diagnostics','/brands/porsche-service-dubai/engine-diagnostics','/services/mercedes-ac-repair-dubai','/brands/bmw-service-dubai/ac-repair','/brands/volkswagen-service-dubai/ac-repair','/brands/range-rover-service-dubai/ac-repair'}

def classify(q):
    n=norm(q); ac=bool(re.search(r'\bac\b|air\s*con|refrigerant|compressor|condenser|evaporator|cooling|مكيف',n))
    brands={'mercedes':'mercedes-benz','bmw':'bmw','porsche':'porsche','volkswagen':'volkswagen','vw':'volkswagen','range rover':'range-rover','cadillac':'cadillac','ferrari':'ferrari','audi':'audi','lamborghini':'lamborghini','rolls royce':'rolls-royce'}
    brand=next((v for k,v in brands.items() if re.search(r'\b'+re.escape(k)+r'\b',n)),None)
    if 'بورش' in n: brand='porsche'
    if 'بي ام دبليو' in n: brand='bmw'
    if brand=='mercedes-benz': owner='/services/mercedes-ac-repair-dubai' if ac else '/services/mercedes-diagnostics-dubai'
    elif brand: owner=f"/brands/{brand}-service-dubai/{'ac-repair' if ac else 'engine-diagnostics'}"
    elif 'electrical' in n or 'electrician' in n: owner='/services/auto-electrical-repair-dubai'
    else: owner='/services/car-ac-repair-dubai' if ac else '/services/car-diagnostics-dubai'
    if re.search(r'[\u0600-\u06ff]',n): owner='/ar'+owner
    intent='Automotive AC inspection/repair enquiry' if ac else 'Diagnostic fault investigation enquiry'
    caveat='Vehicle, system and exact repair coverage must be confirmed before booking; demand does not establish capability.'
    if 'refill' in n: caveat='Refrigerant identification and compatible leak/fault assessment first; a refill is not an automatic remedy.'
    if 'dip2' in n: caveat='DIP2 is a measured query location, not DIGI-TEC premises. Accurate Al Quoz location only; do not create a DIP2 page.'
    if 'ecu' in n or any(t in n for t in ['xentry','ista','piwis','odis']): caveat='Exact vehicle/module function and available platform/access require verification; no universal module repair or programming claim.'
    return owner,intent,caveat

ranked=[
 ('mercedes diagnostic dubai','Strong page-one position6.60 on63impressions; preserve completed diagnostic content and measure actual URL ownership.'),
 ('volkswagen ac repair','65impressions at15.91; existing brand AC owner gives a focused near-page-one repair opportunity.'),
 ('bmw diagnostics dubai','49impressions at13.06; vehicle-specific fault assessment and enquiry clarification fit the existing owner.'),
 ('mercedes ac repair dubai','42impressions,1click and19.86 position; the only recent clicked exact query in this selected cluster.'),
 ('bmw ac repair dubai','38impressions at16.45; targeted diagnosis and cooling-symptom enquiry fit existing owner.'),
 ('porsche diagnostic dubai','28impressions at15.64; reinforce the existing Porsche diagnostics owner without a singular-variant route.'),
 ('range rover ac repair','62impressions at20.05; existing AC owner can explain system and refrigerant limits.'),
 ('fault code scanning dubai','90impressions at25.92; explain code-versus-diagnosis and direct to assessment.'),
 ('car diagnostic dubai','111impressions at28.45; largest recent general diagnostic query that matches the Al Quoz owner.'),
 ('porsche diagnostics dubai','26impressions at18.04; same owner as singular query and no additional page.'),
 ('engine diagnostics dubai','62impressions at24.66; repair-led investigation intent on the general diagnostics owner.'),
 ('car ac gas refill dubai','220impressions at35.88; greatest visible AC demand, but scope requires leak/fault assessment rather than a refill promise.'),
 ('car ac repair dubai','61impressions at38.61; direct commercial service intent, broader and more competitive than brand AC terms.'),
 ('auto electrical diagnostics dubai','82impressions at36.72; preserve strong existing electrical owner and connect diagnostics journey.'),
 ('vehicle diagnostics dubai','74impressions at34.99; same general diagnostic owner, natural coverage and booking path.'),
 ('car diagnostic service dubai','64impressions at37.16; general service assessment variant on existing diagnostics owner.'),
 ('porsche engine diagnostics','35impressions at19.60; focused brand-owner relevance, without inventing exact tool compatibility.'),
 ('car ac refill services dubai','71impressions at27.90; retain on general AC owner with diagnosis and refrigerant constraints.'),
 ('volkswagen ac repair dubai','19impressions at14.74; explicit Dubai variant groups with the stronger non-city VW query.'),
 ('mercedes benz ac repair','35impressions at20.29 plus5fresh24himpressions; same Mercedes AC owner.')
]
rankmap={norm(q):(i+1,why) for i,(q,why) in enumerate(ranked)}
measured=[r['query'] for r in evidence if r['source_file']==PRIMARY['92d']]
for q in ['porsche ac repair dubai','bmw engine diagnostics','range rover diagnostic dubai','تصليح مكيف بي ام دبليو']:
    if norm(q) not in {norm(x) for x in measured}: measured.append(q)
proposed=['Mercedes XENTRY diagnostics Dubai','BMW drivetrain malfunction Dubai','VW EPC light Dubai','Check engine light Dubai','Car electrical diagnostics Dubai','Car AC not cooling Dubai','Car AC gas leak repair Dubai','Car AC compressor replacement Dubai','Car AC condenser replacement Dubai','Car AC evaporator repair Dubai']
rows=[]
for q in measured+proposed:
    owner,intent,caveat=classify(q); record=owners.get(owner); ar=owner if owner.startswith('/ar/') else '/ar'+owner
    measured_class= q in measured
    row={'query':q,'evidence_class':'Measured exact raw GSC query' if measured_class else 'Proposed variant; no exact row found in inspected visible exports','intent':intent,'structural_owner':owner if record else '', 'structural_owner_basis':'Verified prior release initial HTML, canonical, sitemap and locale counterpart audit; editorial intent assignment only' if record else 'Provisional; owner requires route audit','owner_indexable':bool(record and not any('noindex' in x for x in record['robots'])),'owner_in_sitemap':record['inSitemap'] if record else None,'arabic_owner':ar if ar in owners else '', 'gsc_query_by_page_owner':None,'gsc_query_by_page_status':'Unavailable: Queries and Pages are separate dimensions; no exact diagnostics/AC query-filtered Pages export or joined API response supplied','commercial_relevance':'High when a supported vehicle needs diagnosis or repair','capability_boundary':caveat,'top20_rank':rankmap.get(norm(q),(None,None))[0],'ranking_reason':rankmap.get(norm(q),(None,None))[1] or 'Supporting variant or backlog; not a separate page opportunity','priority':'P1' if owner in focus and norm(q) in rankmap else 'P2'}
    if not measured_class: row['priority']='P2: PROPOSED / VERIFY'
    if 'dip2' in norm(q): row['priority']='HOLD: LOCATION INTENT';row['commercial_relevance']='Commercial, but location mismatch requires truthful Al Quoz context'
    if 'ecu diagnostic' in norm(q) or 'xentry' in norm(q): row['priority']='HOLD: CAPABILITY';row['commercial_relevance']='Commercial only if exact requested function is supported'
    if any(x in norm(q) for x in ['auto air con','car diagnostic']) and norm(q) in ['auto air con','car diagnostic']: row['priority']='P3: BROAD / AMBIGUOUS';row['commercial_relevance']='Mixed general intent; automotive and Dubai qualification needed'
    if owner=='/services/mercedes-diagnostics-dubai': row['recommended_action']='Preserve completed repair-sprint diagnostic content; monitor and validate query-by-page ownership. No repeat hub or diagnostic rewrite.'
    elif owner=='/services/auto-electrical-repair-dubai': row['recommended_action']='Preserve completed electrical repair content; link the diagnostic assessment journey where relevant and monitor.'
    elif owner in focus: row['recommended_action']='Improve existing owner with symptom-to-test explanation, exact-vehicle scope and service-specific assessment enquiry; retain established routes and entity IDs.'
    else: row['recommended_action']='Backlog for existing owner; verify exact capability and query-by-page data before expanding. No new route.'
    if 'dip2' in norm(q): row['recommended_action']='Retain measured evidence as location-specific backlog; no DIP2 premises claim or new location page.'
    if not measured_class: row['recommended_action']='Proposed supporting FAQ/symptom coverage on existing owner only after capability verification; no demand or ranking claim.'
    for label,name in PRIMARY.items():
        match=lookups[label].get(norm(q)); src=books[name]['source']; row[label+'_observed']=bool(match)
        row[label+'_clicks']=match[1]['Clicks'] if match else None;row[label+'_impressions']=match[1]['Impressions'] if match else None;row[label+'_ctr']=match[1]['CTR'] if match else None;row[label+'_position']=match[1]['Position'] if match else None
        row[label+'_source_file']=name;row[label+'_source_period']=f"{src['start']} to {src['end']}";row[label+'_source_time_column']=src['time_column'];row[label+'_filters']=json.dumps(src['filters'],ensure_ascii=False,separators=(',',':'));row[label+'_source_range']=f'Queries!A{match[0]}:E{match[0]}' if match else None;row[label+'_missing_meaning']='' if match else 'Absent from visible exported query rows; not zero demand'
    all_matches=[e for e in evidence if norm(e['query'])==norm(q)]
    row['all_period_observation_count_not_summed']=len(all_matches)
    if not measured_class: assert not any(norm(x['Top queries'])==norm(q) for b in books.values() for x in b['data']['Queries']), f'Proposed keyword is measured: {q}'
    row['competition']='Dedicated diagnostic and AC workshop competitors observed on primary business pages; qualitative review, not numeric keyword difficulty or a localized rank report.'
    row['competition_review_source']='qualitative-competition-review.json'
    if row['ranking_reason']: row['ranking_reason']=readable(row['ranking_reason'])
    rows.append(row)
rows.sort(key=lambda r:(r['top20_rank'] is None,r['top20_rank'] or 999,r['evidence_class'],r['query']))
write_json('keyword-opportunity-matrix.json',rows);write_csv('keyword-opportunity-matrix.csv',rows)
top20=[r for r in rows if r['top20_rank']]; assert len(top20)==20
write_json('top20-commercial-opportunities.json',top20);write_csv('top20-commercial-opportunities.csv',top20)
write_csv('next-keyword-backlog.csv',[r for r in rows if not r['top20_rank']])

attachment=Path(r'C:\Users\ADMIN\.codex\attachments\102abc1d-898f-43af-8c41-09e792cc85fa\Pasted text.txt').read_text(encoding='utf-8')
section=attachment.split('## A2. Diagnostics keyword research')[1].split('# PHASE B')[0]
requested=[line[2:].strip() for line in section.splitlines() if line.startswith('- ')]
requested_rows=[]
for q in requested:
    for label,name in PRIMARY.items():
        match=lookups[label].get(norm(q));src=books[name]['source']
        requested_rows.append({'research_candidate':q,'source_file':name,'source_period':f"{src['start']} to {src['end']}",'source_time_column':src['time_column'],'source_filters':json.dumps(src['filters'],ensure_ascii=False,separators=(',',':')),'exact_query_observed':bool(match),'exact_source_query':match[1]['Top queries'] if match else None,'clicks':match[1]['Clicks'] if match else None,'impressions':match[1]['Impressions'] if match else None,'ctr':match[1]['CTR'] if match else None,'position':match[1]['Position'] if match else None,'source_range':f'Queries!A{match[0]}:E{match[0]}' if match else None,'missing_meaning':'' if match else 'Not present in visible export; not evidence of zero demand'})
write_csv('requested-candidate-baselines.csv',requested_rows)

competition=[
 {'cluster':'General diagnostics','assessment':'Strong visible specialist competition; dedicated diagnostic pages, scan-versus-diagnosis explanations and direct booking paths. Qualitative only, not keyword difficulty or a local rank audit.','urls':['https://activeauto.me/services/vehicle-diagnostics','https://aljubailgarage.com/services/computer-diagnostics','https://www.fixgo.ae/services/car-computer-diagnostics-dubai']},
 {'cluster':'General AC','assessment':'Strong visible dedicated AC competition with symptom, leak and compressor explanations and immediate enquiry options. Competitor service and price claims are not evidence of DIGI-TEC capability.','urls':['https://caracheroes.com/','https://nkauto-me.com/services/ac-repair/']},
 {'cluster':'Mercedes/BMW diagnostics','assessment':'Specialist brand-focused diagnostic and electrical fault pages visible. Support vehicle-specific fault assessment and clear limitations; no need to copy capability claims.','urls':['https://mbcod.com/']},
 {'cluster':'BMW AC','assessment':'Multiple dedicated brand AC pages visible, including symptom-led process and quotation paths. High commercial relevance but no guaranteed ranking gain.','urls':['https://bmwgarage.ae/ac/','https://kartalauto.ae/ac-repair/bmw/','https://www.dascenter.ae/bmw-ac-repair-dubai/']},
 {'cluster':'Volkswagen AC','assessment':'Dedicated Volkswagen AC competitors visible with diagnosis/process content. Existing owner and position14.74–15.91 support focused relevance and conversion work.','urls':['https://volkswagenworkshopdubai.com/volkswagen-ac-repair-dubai/','https://euroexpert.ae/service/volkswagen-ac-repair-dubai/']},
 {'cluster':'Porsche diagnostics','assessment':'Dedicated independent Porsche diagnostic page visible with warning and system intent. Preserve exact-vehicle limitations and differentiate using verified process.','urls':['https://porsche-repair-dubai.com/services/porsche-diagnostics/']}
]
write_json('qualitative-competition-review.json',{'reviewed_at':'2026-10-19','method':'Web search results and primary business page contents, unpersonalized tool location unspecified. No numeric keyword difficulty, complete SERP order, paid volume or local map ranking claimed.','clusters':competition})
checks={'source_workbooks':len(manifest),'all_raw_sources_unchanged_sha256':all(x['source_unchanged'] for x in manifest),'opportunities':len(rows),'measured_distinct_queries':len(measured),'proposed_distinct_queries':len(proposed),'measured_recent_92d':len([r for r in rows if r['92d_observed']]),'measured_historical_only_or_24h':len([r for r in rows if r['evidence_class'].startswith('Measured') and not r['92d_observed']]),'top20':len(top20),'all_top20_measured_92d':all(r['92d_observed'] for r in top20),'all_owners_verified':all(r['structural_owner'] for r in rows),'all_owners_indexable':all(r['owner_indexable'] for r in rows),'all_owners_in_sitemap':all(r['owner_in_sitemap'] for r in rows),'all_have_arabic_owner':all(r['arabic_owner'] for r in rows),'query_by_page_evidence_available':False,'overlapping_periods_combined':False,'missing_metrics_coerced_to_zero':False,'requested_research_candidates':len(requested),'raw_query_observations_all_periods':len(evidence),'source_semantic_duplicates':[{ 'file':Path(s['path']).name,'duplicate_of':s['semantic_duplicate_of']} for s in manifest if s['semantic_duplicate_of']]}
assert 30<=len(rows)<=50;assert checks['all_top20_measured_92d'];assert checks['all_owners_verified'];assert checks['all_owners_indexable'];assert checks['all_owners_in_sitemap'];assert checks['all_have_arabic_owner']
write_json('analysis-checks.json',checks)
lines=['DIGI-TEC DIAGNOSTICS AND AC GSC RESEARCH — 19 OCTOBER 2026','',f"{len(rows)} opportunities: {len(measured)} distinct measured raw queries and {len(proposed)} proposed variants. {checks['measured_recent_92d']} measured queries occur in the92-day export; the remainder are labelled historical/24-hour evidence. Top20 all have current92-day evidence.",'','Source basis:','- Primary92-day:10 July–9 October2026, Web, no query/page/country/device filter in workbook.','- Historical28-day:6 August–2 September2026, Web, no query/page filter. This overlaps the92-day period and is not summed or used as an independent before/after window.','- Last24hours export19October:18October06:00 to19October05:00, source time-column label retained in manifest; Web(text-based). Small fresh sample, not an exact daily trend or a post-implementation result.','-17original workbooks inspected read-only and SHA256 checked unchanged. Historical7-day, older92-day,167-day, query-filtered Mercedes, page-filtered tuning, PPF-filtered and Product-snippets exports are retained with filters; none pooled with primary periods.','- Some1000-row query tables are capped, and even shorter exports omit anonymized or non-exported queries. Missing means unavailable, never measured zero demand.','- CTR retained exactly as exported(as ratio); clicks/impressions/average position kept at exact query/source grain. No sum of periods, queries or average positions used to claim total demand.','- No exact diagnostics/AC query-by-page joined export/API response supplied. Existing route assignment is verified structural/editorial ownership, not a GSC ranking-page observation. Independent Pages rows are saved separately and never joined by row order or keyword similarity.','','Priority method: analyst ordering considers primary92-day impressions, position(especially6–20), explicit automotive repair/Dubai intent, existing indexable owner and capability/location constraints. Qualitative competition is recorded from primary business sites. There is no invented monthly volume, numeric difficulty, conversion rate, lead count or revenue estimate. Similar variants share one existing owner.','', 'Implementation selection: general diagnostics/generalAC; BMW/Porsche diagnostics; Mercedes/BMW/Volkswagen/RangeRoverAC. Completed Mercedes diagnostics and electrical repair copy are preserved. PorscheAC and other brands remain existing-owner backlog. No new route is justified by these exports.','','Top20(query:92-day clicks/impressions/CTR/position):']
for r in top20: lines.append(f"{r['top20_rank']}. {r['query']}: {r['92d_clicks']:g}/{r['92d_impressions']:g}/{r['92d_ctr']:.2%}/{r['92d_position']:.2f}. {r['structural_owner']}. {r['ranking_reason']}")
lines+=['','Capability and location boundaries:','- DIP2 queries are measured but do not authorize aDIP2 location/service claim. DIGI-TEC location remains verified Al Quoz.','- Named-tool variants(XENTRY/ISTA+/PIWIS/ODIS/JLR), ECU programming, compressor/condenser/evaporator replacement and high-voltage work require exact workshop and vehicle evidence. Query demand and existing marketing text do not verify this independently.','- Refill demand should be answered with refrigerant identification, diagnosis and applicable leak checks before quotation, without inventing refrigerant compatibility or refill prices.','- Warning-light, no-start, misfire and warm-air queries can support factual symptom FAQs on existing owners. Proposed variants have blank metrics and are not counted as measured demand.','- WhatsApp/tel interactions are contact attempts, not confirmed enquiries. Closed-loop qualified enquiries require actual received lead records and consent-aware attribution.','','Needed next GSC evidence: export Search Analytics with dimensions query+page for these measured clusters, dates and country/device/search appearance filters retained; exact-query URL Inspection for canonical/indexing checks; comparable nonoverlapping28-day query exports after release. The daily Chart sheet alone cannot produce query-level28-day metrics.','','Primary competition sources:']
for c in competition:
    lines.append(c['cluster']+': '+c['assessment']);lines.extend('  '+u for u in c['urls'])
report=readable('\n'.join(lines)).replace('to19 October','to 19 October').replace('Web(text-based)','Web (text-based)').replace('column label retained in manifest','column is Time (UTC+04:00), Dubai time').replace('day:10','day: 10').replace('day:6','day: 6').replace('October:18','October: 18').replace('day,167','day, 167').replace('Top 20(query:92','Top 20 (query: 92').replace('Strong page-one average','Strong average')
report+='\nCommercial rank and implementation priority differ: an existing completed owner may rank highly for monitoring while its content remains preserved. No GSC average position is a guarantee of any particular live search result.\n'
(OUT/'GSC-RESEARCH-REPORT.txt').write_text(report+'\n',encoding='utf-8')
print(json.dumps(checks,indent=2,ensure_ascii=False))
