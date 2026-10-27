"""Read-only source extraction and five-brand GSC opportunity analysis."""
from pathlib import Path
import csv
import hashlib
import json
import re
import unicodedata
import warnings
import openpyxl

OUT = Path(__file__).resolve().parent
SOURCES = Path(r'C:\Users\ADMIN\Downloads')
warnings.filterwarnings('ignore', message='Workbook contains no default style')

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def total(rows):
    clicks = sum(float(r.get('Clicks') or 0) for r in rows)
    impressions = sum(float(r.get('Impressions') or 0) for r in rows)
    position = sum(float(r.get('Position') or 0) * float(r.get('Impressions') or 0) for r in rows) / impressions if impressions else None
    return {'clicks': clicks, 'impressions': impressions, 'ctr': clicks / impressions if impressions else None, 'weighted_position': position}

manifest, books = [], {}
for source in sorted(SOURCES.glob('digitecme.com-Performance-on-Search-2026-*.xlsx')):
    before = digest(source)
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheets, ranges = {}, {}
    for sheet in book:
        values = list(sheet.values)
        headers = list(values[0]) if values else []
        sheets[sheet.title] = [dict(zip(headers, row)) for row in values[1:] if any(v is not None for v in row)]
        ranges[sheet.title] = f'A1:{openpyxl.utils.get_column_letter(sheet.max_column)}{sheet.max_row}'
    book.close()
    assert before == digest(source), f'Source modified: {source}'
    filters = {r.get('Filter'): r.get('Value') for r in sheets.get('Filters', [])}
    date_key = next((key for key in (sheets.get('Chart') or [{}])[0] if key == 'Date' or key.startswith('Time (')), None)
    dates = [str(r[date_key]) for r in sheets.get('Chart', []) if date_key and r.get(date_key)]
    item = {'path': str(source), 'sha256': before, 'filters': filters, 'time_column': date_key, 'start': min(dates) if dates else None, 'end': max(dates) if dates else None, 'time_rows': len(dates), 'sheets': {name: {'range': ranges[name], 'rows': len(rows), **(total(rows) if name != 'Filters' else {})} for name, rows in sheets.items()}}
    manifest.append(item)
    books[source.name] = {'source': item, 'data': sheets}

OUT.mkdir(parents=True, exist_ok=True)
(OUT / 'source-manifest.json').write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding='utf-8')
(OUT / 'extracted-workbooks.json').write_text(json.dumps(books, indent=2, ensure_ascii=False), encoding='utf-8')

PRIMARY = [('baseline_24h', '2026-10-19'), ('recent_92d', '2026-10-12 (1)'), ('historical_167d', '2026-09-27')]
brands = {'Rolls-Royce': 'rolls-royce', 'Lamborghini': 'lamborghini', 'Volkswagen': 'volkswagen', 'Mercedes-Benz': 'mercedes-benz', 'BMW': 'bmw'}
brand_re = {'Rolls-Royce': r'\brolls\s*royce\b|رولز|رولس', 'Lamborghini': r'\blamborghini\b|لامبور', 'Volkswagen': r'\bvolkswagen\b|\bvw\b|فولكس', 'Mercedes-Benz': r'\bmercedes\b|مرسيدس', 'BMW': r'\bbmw\b|بي ام دبليو|بي إم دبليو'}
targets = [(brand, f'{("mercedes" if brand == "Mercedes-Benz" else brand.lower().replace("-", " "))} {kind} dubai') for brand in brands for kind in ('service', 'repair')]

def norm(text):
    return re.sub(r'\s+', ' ', re.sub(r'[-–—]', ' ', unicodedata.normalize('NFKC', text).lower())).strip()

owner_records = {}
root = OUT.parents[2]
for rel, field in [('outputs/b5/rolls-royce-keyword-owner-records.csv', 'raw_keyword'), ('outputs/b4/lamborghini-keyword-owner-records.csv', 'raw_keyword'), ('outputs/b8/volkswagen-keyword-owner-records.csv', 'raw_keyword'), ('outputs/b3/bmw-keyword-owner-records.csv', 'raw_keyword'), ('outputs/b1/keyword-owner-records.csv', 'keyword_raw')]:
    with (root / rel).open(encoding='utf-8-sig', newline='') as stream:
        for record in csv.DictReader(stream):
            key = norm(record.get(field, ''))
            if key and record.get('primary_owner'):
                owner_records.setdefault(key, set()).add(record['primary_owner'])

def ownership(query, brand):
    q = norm(query)
    owners = owner_records.get(q, set())
    if len(owners) == 1:
        return next(iter(owners)), 'Established editorial ownership register'
    if len(owners) > 1:
        return '; '.join(sorted(owners)), 'Multiple historical register owners: review before targeting'
    if re.search(r'[\u0600-\u06ff]', q):
        return f'/ar/brands/{brands[brand]}-service-dubai', 'Existing Arabic counterpart; editorial intent grouping, not measured ranking URL'
    if re.search(r'\btun(?:ing|ning|e)\b|\bstage\b|\bremap\b|\bchip\b', q):
        return '/tuning', 'Tuning intent outside service sprint; existing owner only, verify capability before targeting'
    if re.search(r'\b(?:audio|stereo|speaker|sound system)\b', q):
        return '/services/mercedes-audio-upgrade-dubai' if brand == 'Mercedes-Benz' else '', 'Specific audio task: preserve existing task owner if verified'
    if re.search(r'vin decoder|\bforsale\b|\bfor sale\b|\bjobs\b|\bsalary\b|\btraining\b|\bcourse\b', q):
        return '', 'Different intent; no service landing page targeting justified'
    return f'/brands/{brands[brand]}-service-dubai', 'Provisional hub grouping; exact specialized task/model owner needs route and capability review'

query_rows, target_rows, page_rows, device_rows, family_rows = [], [], [], [], []
for label, suffix in PRIMARY:
    name = f'digitecme.com-Performance-on-Search-{suffix}.xlsx'
    book = books[name]
    source = book['source']
    period = source['start'] + ' to ' + source['end']
    query_lookup = {norm(row['Top queries']): (i, row) for i, row in enumerate(book['data']['Queries'], 2)}
    for brand, query in targets:
        match = query_lookup.get(norm(query))
        rownum, row = match if match else (None, {})
        target_rows.append({'period_label': label, 'source_file': name, 'source_period': period, 'brand': brand, 'query': query, 'observed': bool(match), 'clicks': row.get('Clicks'), 'impressions': row.get('Impressions'), 'ctr': row.get('CTR'), 'position': row.get('Position'), 'source_range': f'Queries!A{rownum}:E{rownum}' if match else '', 'editorial_owner': f'/brands/{brands[brand]}-service-dubai', 'observed_ranking_page': 'Not available from separate Queries/Pages tabs'})
    for rownum, row in enumerate(book['data']['Queries'], 2):
        query = row['Top queries']
        matched = [brand for brand, regex in brand_re.items() if re.search(regex, norm(query))]
        for brand in matched:
            owner, evidence = ownership(query, brand)
            query_rows.append({'period_label': label, 'source_file': name, 'source_period': period, 'brand': brand, 'query': query, 'clicks': row['Clicks'], 'impressions': row['Impressions'], 'ctr': row['CTR'], 'position': row['Position'], 'editorial_owner': owner, 'ownership_evidence': evidence, 'source_range': f'Queries!A{rownum}:E{rownum}', 'observed_ranking_page': 'Unavailable; no query × page dimension'})
    for rownum, row in enumerate(book['data']['Pages'], 2):
        for brand, slug in brands.items():
            if f'/brands/{slug}-service-dubai' in row['Top pages'] or (brand == 'Mercedes-Benz' and '/mercedes/' in row['Top pages']):
                page_rows.append({'period_label': label, 'source_file': name, 'source_period': period, 'brand': brand, 'url': row['Top pages'], 'clicks': row['Clicks'], 'impressions': row['Impressions'], 'ctr': row['CTR'], 'position': row['Position'], 'source_range': f'Pages!A{rownum}:E{rownum}', 'query_attribution': 'Not available from this page row'})
    for rownum, row in enumerate(book['data']['Devices'], 2):
        device_rows.append({'period_label': label, 'source_file': name, 'source_period': period, **row, 'source_range': f'Devices!A{rownum}:E{rownum}', 'scope': 'Sitewide; cannot attribute to five brands'})
    for brand in brands:
        selected = [r for r in query_rows if r['period_label'] == label and r['brand'] == brand]
        family_rows.append({'period_label': label, 'source_file': name, 'brand': brand, 'source_period': period, 'visible_query_rows': len(selected), 'clicks': sum(r['clicks'] for r in selected), 'impressions': sum(r['impressions'] for r in selected), 'scope': 'Visible explicit-brand query rows, includes specialized and adjacent intents; not all brand demand'})

def csvout(name, rows):
    with (OUT / name).open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]) if rows else [])
        writer.writeheader()
        writer.writerows(rows)

csvout('target-query-baselines.csv', target_rows)
csvout('five-brand-query-register.csv', query_rows)
csvout('five-brand-page-baselines.csv', page_rows)
csvout('sitewide-device-baselines.csv', device_rows)
csvout('visible-brand-query-summary.csv', family_rows)

priority_selection = {
    'mercedes repair dubai': ('P1', 'Improve hub snippet relevance and first-screen booking clarity; preserve established repair owner', 'Highest recent exact target-query impressions with average position 6.35 and CTR 0.75%; CTR opportunity requires review, not a guaranteed gain'),
    'mercedes service dubai': ('P1', 'Make routine maintenance scope and quote process easy to find on the same hub', '613 impressions and average position 13.88 support relevance and internal-link work'),
    'rolls royce service dubai': ('P1', 'Strengthen verified service coverage and booking process on the established Rolls-Royce hub', '476 recent impressions and average position 18.13; 12 fresh baseline impressions at 11.08'),
    'rolls royce repair dubai': ('P1', 'Describe verified repair assessment and link established mechanical, diagnostic and electrical owners', '354 impressions, 0.28% CTR, position 19.92; same broad hub as service'),
    'lamborghini service dubai': ('P1', 'Clarify maintenance, diagnostics, service assessment and quote steps on the existing hub', '379 impressions and no clicks at average position 16.41'),
    'lamborghini repair dubai': ('P1', 'Use existing hub for broad repair intent and link relevant repair children', '361 impressions and no clicks at average position 22.37'),
    'bmw service dubai': ('P1', 'Strengthen hub maintenance relevance and supported model/service paths', '377 impressions, one click, average position 21.52; missing from 24h table does not remove opportunity'),
    'bmw repair dubai': ('P1', 'Clarify diagnostics-led repair and enquiry process on the same BMW hub', '203 impressions, two clicks, average position 13.79; baseline average 7.25 on four impressions'),
    'volkswagen service dubai': ('P2', 'Add natural Volkswagen and VW language, verified service scope and clear enquiry next step', '107 impressions, no clicks, average position 22.48; relative volume is lower but commercially relevant'),
    'volkswagen repair dubai': ('P2', 'Retain the existing Volkswagen hub for broad repair; make task routes easy to follow', '89 impressions, no clicks, average position 26.53'),
    'mercedes specialist dubai': ('P1', 'Improve independent-workshop relevance with verified facts; avoid certification/dealer implications', '341 impressions, no clicks, average position 5.8'),
    'mercedes workshop dubai': ('P1', 'Use existing hub and accurate Al Quoz location/booking information', '422 impressions, no clicks, average position 15.96'),
    'mercedes service center dubai': ('P1', 'Keep broad independent-service intent on the Mercedes hub', '353 impressions, no clicks, average position 12.76'),
    'rolls royce workshop dubai': ('P2', 'Cover independent-workshop selection naturally on hub and link existing supporting guide', '148 impressions, no clicks, average position 18.54'),
    'rolls royce service center dubai': ('P2', 'Use accurate independent workshop identity and one broad service owner', '138 impressions, no clicks, average position 17.49'),
    'lamborghini service center dubai': ('P2', 'Keep broad service-center searches on the hub without dealer status implications', '195 impressions, no clicks, average position 20.19'),
    'lamborghini workshop dubai': ('P2', 'Improve existing hub and supporting selection-guide links', '165 impressions, no clicks, average position 15.44'),
    'bmw service center dubai': ('P2', 'Improve existing hub scope and navigation without creating a duplicate keyword page', '188 impressions, no clicks, average position 30.43'),
    'bmw workshop dubai al quoz': ('P2', 'Reinforce verified workshop location and directions on existing hub', '56 impressions, no clicks, average position 14.09'),
    'vw service dubai': ('P2', 'Include VW as a natural brand abbreviation within the same Volkswagen hub', '32 impressions, no clicks, average position 25.41'),
    'volkswagen ac repair': ('P2', 'Link existing AC repair child from hub; keep detailed AC work on that owner', '65 impressions, no clicks, average position 15.91'),
    'bmw transmission repair dubai': ('P2', 'Link existing transmission child and keep capability claims conditional on assessment', '90 impressions, no clicks, average position 11.6'),
    'bmw engine repair dubai': ('P2', 'Link mechanical-repair child from hub using descriptive engine/mechanical context', '74 impressions, no clicks, average position 16.35'),
    'mercedes engine repair dubai': ('P2', 'Preserve established mechanical-repair service owner and use hub as a routing page', '271 impressions, no clicks, average position 10.78'),
    'mercedes suspension repair dubai': ('P2', 'Preserve established suspension service owner and link from hub', '235 impressions, no clicks, average position 13.6'),
    'rolls royce engine repair dubai': ('P2', 'Link existing mechanical-repair child with assessment-led wording', '115 impressions, no clicks, average position 26.01'),
    'rolls royce suspension repair dubai': ('P2', 'Retain suspension intent on existing service child', '66 impressions, no clicks, average position 19.33'),
    'lamborghini engine repair dubai': ('P2', 'Route detailed engine/mechanical repair to existing child', '70 impressions, no clicks, average position 25.16'),
    'lamborghini transmission repair dubai': ('P2', 'Route detailed transmission repair to existing child', '49 impressions, no clicks, average position 15.14'),
    'lamborghini battery replacement dubai': ('P2', 'Link existing battery service child and explain assessment before quotation', '53 impressions, no clicks, average position 27.32'),
    'lamborghini reverse camera in dubai': ('HOLD: VERIFY', 'Review existing electrical owner for diagnostic/repair scope before writing any camera claim', '247 impressions are observed demand; they do not prove reverse-camera installation or repair capability'),
    'bmw coding service': ('HOLD: VERIFY', 'Preserve established diagnostic owner; verify exact task and vehicle support before targeting', '88 impressions at position 38.77 do not verify coding compatibility'),
    'volkswagen vin decoder dubai': ('EXCLUDE', 'Do not optimize commercial workshop pages for a decoder-tool request', '32 impressions represent a distinct tool intent'),
    'صيانة رولز رويس': ('P2', 'Improve existing Arabic hub with matching verified content and enquiry process', '43 impressions, no clicks, average position 28.93'),
}
priorities = []
for row in query_rows:
    if row['period_label'] != 'recent_92d' or norm(row['query']) not in priority_selection:
        continue
    priority, action, why = priority_selection[norm(row['query'])]
    priorities.append({'priority': priority, **row, 'action': action, 'reason': why, 'implementation_guardrail': 'Existing owner only; no new per-keyword page, capability/price/certification/warranty claims or inferred conversion counts'})
priorities.sort(key=lambda r: (r['priority'], -r['impressions']))
csvout('prioritized-keyword-opportunities.csv', priorities)

recent = books['digitecme.com-Performance-on-Search-2026-10-12 (1).xlsx']['data']['Chart']
comparison = []
for start, end in [('2026-08-15', '2026-09-11'), ('2026-09-12', '2026-10-09')]:
    rows = [r for r in recent if start <= r['Date'] <= end]
    assert len(rows) == 28
    comparison.append({'source': 'digitecme.com-Performance-on-Search-2026-10-12 (1).xlsx', 'start': start, 'end': end, 'days': len(rows), **total(rows), 'scope': 'Sitewide Chart daily rows; same source and nonoverlapping windows; not five-brand performance'})
(OUT / 'nonoverlapping-sitewide-comparison.json').write_text(json.dumps(comparison, indent=2), encoding='utf-8')

assert len(target_rows) == 30
assert len([r for r in target_rows if r['period_label'] == 'baseline_24h' and r['observed']]) == 9
assert all(r['observed'] for r in target_rows if r['period_label'] == 'recent_92d')
assert all(len(b['data']) == 7 for b in books.values())
(OUT / 'analysis-checks.json').write_text(json.dumps({'source_workbook_count': len(books), 'all_sources_unchanged_sha256': True, 'seven_sheets_inspected_per_workbook': True, 'primary_source_count': len(PRIMARY), 'target_rows': len(target_rows), 'recent_92d_exact_targets_observed': 10, 'baseline_24h_exact_targets_observed': 9, 'baseline_missing_exact_target': 'bmw service dubai', 'priority_records': len(priorities), 'same_source_nonoverlapping_chart_comparison': True}, indent=2), encoding='utf-8')

print(json.dumps({'baseline': books['digitecme.com-Performance-on-Search-2026-10-19.xlsx']['source'], 'target_queries': target_rows[:20], 'brand_summaries': family_rows, 'chart_comparison': comparison}, indent=2, ensure_ascii=False))
