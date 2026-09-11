from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urlsplit, unquote
import csv, html, json, re, hashlib
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs/mercedes-release-2026-09-08'
OUT.mkdir(exist_ok=True)
PLAN = json.loads((Path(__file__).parent / 'plan.json').read_text(encoding='utf-8'))

class Document(HTMLParser):
    def __init__(self, text):
        super().__init__(convert_charrefs=True)
        self.meta = {}; self.links = []; self.assets = []; self.ids = set()
        self.feed(text)
        self.title = self.extract(text, 'title')
        self.h1 = [self.clean(x) for x in re.findall(r'<h1\b[^>]*>(.*?)</h1>', text, re.S)]
        self.headings = [self.clean(x) for x in re.findall(r'<h[23]\b[^>]*>(.*?)</h[23]>', text, re.S)]
        self.schemas = [json.loads(x) for x in re.findall(r'<script[^>]*type="application/ld\+json"[^>]*>(.*?)</script>', text, re.S)]
        self.visible_source = self.clean(re.sub(r'<(script|style)\b[^>]*>.*?</\1>', '', text, flags=re.S))
    @staticmethod
    def clean(x): return ' '.join(html.unescape(re.sub(r'<[^>]*>', ' ', x)).split())
    def extract(self, text, tag):
        found = re.findall(fr'<{tag}\b[^>]*>(.*?)</{tag}>', text, re.S)
        assert len(found) == 1, f'{tag} count {len(found)}'
        return self.clean(found[0])
    def handle_starttag(self, tag, attrs):
        d = dict(attrs)
        if d.get('id'): self.ids.add(d['id'])
        if tag == 'meta': self.meta[d.get('name') or d.get('property')] = d.get('content')
        if tag == 'a' and d.get('href'): self.links.append(d['href'])
        if tag == 'link' and d.get('rel') == 'canonical': self.meta['canonical'] = d.get('href')
        if tag == 'link' and d.get('rel') == 'alternate': self.meta.setdefault('alternates', []).append(d)
        if tag in ['img','script','source'] and d.get('src'): self.assets.append(d['src'])
        if tag == 'link' and d.get('rel') == 'stylesheet': self.assets.append(d['href'])

def html_file(path): return ROOT / 'dist' / path.strip('/') / 'index.html'

def inspect(path, indexable=True):
    doc = Document(html_file(path).read_text(encoding='utf-8'))
    assert len(doc.h1) == 1, path
    assert doc.meta.get('canonical') == 'https://digitecme.com' + path, path
    assert doc.meta.get('description'), path
    assert doc.meta.get('robots','').startswith('index,' if indexable else 'noindex,'), path
    assert ('https://digitecme.com' + path in sitemap) == indexable, path
    checked_links = 0
    for link in doc.links:
        u = urlsplit(link)
        if u.scheme and u.scheme not in ['http','https']: continue
        if u.netloc and u.netloc not in ['digitecme.com','www.digitecme.com']: continue
        if not u.path.startswith('/'): continue
        target = unquote(u.path)
        if target in redirects: continue
        file = html_file(target)
        assert file.is_file() or (ROOT/'dist'/target.strip('/')).is_file(), f'{path}: missing link {link}'
        if u.fragment and file.is_file():
            assert unquote(u.fragment) in Document(file.read_text(encoding='utf-8')).ids, f'{path}: missing anchor {link}'
        checked_links += 1
    for asset in doc.assets:
        u = urlsplit(asset)
        if u.netloc or not u.path.startswith('/'): continue
        assert (ROOT/'dist'/unquote(u.path).strip('/')).is_file(), f'{path}: missing asset {asset}'
    nodes=[]
    for schema in doc.schemas: nodes += schema.get('@graph',[schema])
    faq_nodes=[n for n in nodes if n.get('@type')=='FAQPage']
    for faq in faq_nodes:
        for q in faq.get('mainEntity',[]):
            assert q['name'] in doc.visible_source, f'{path}: FAQ question not rendered: {q["name"]}'
    return {'path':path,'title':doc.title,'description':doc.meta['description'],'h1':doc.h1[0],
        'canonical':doc.meta['canonical'],'robots':doc.meta['robots'],'headings':doc.headings,
        'schemaTypes':[n.get('@type') for n in nodes], 'checkedInternalLinks':checked_links,
        'checkedAssets':len(doc.assets),'sha256':hashlib.sha256(html_file(path).read_bytes()).hexdigest(),
        'alternates':doc.meta.get('alternates',[])}

sitemap=(ROOT/'dist/sitemap.xml').read_text(encoding='utf-8')
redirects={line.split()[0] for line in (ROOT/'dist/_redirects').read_text().splitlines() if line and not line.startswith('#')}
assert len(PLAN['queries']) == 353 and len(PLAN['pages']) == 31
assert len({q['query'] for q in PLAN['queries']}) == 353
assert len({p['id'] for p in PLAN['pages']}) == 31
ids={p['id'] for p in PLAN['pages']}
assert all(q['id'] in ids for q in PLAN['queries'])

# Read the requested four workbook tabs and reconcile every original measure.
wb=openpyxl.load_workbook(Path(__file__).parent/'Mercedes-keyword-page-plan-2026-09-08.xlsx',data_only=True,read_only=True)
tabs={name:list(wb[name].values) for name in ['Query map','Page priorities','Content briefs','Live checks']}
original=openpyxl.load_workbook(Path('C:/Users/ADMIN/Downloads/digitecme.com-Performance-on-Search-2026-09-08 (1).xlsx'),data_only=True,read_only=True)
raw=[r[:5] for r in list(original['Queries'].values)[1:] if r[0]]
mapped=[r[:5] for r in tabs['Query map'] if r[0] and isinstance(r[1], (int,float))]
assert len(raw)==353 and mapped==raw, 'Original Search Console measures were changed'

records=[]
for row in PLAN['pages']:
    if row['id']=='AUDIO':
        assert not html_file(row['path']).exists(), 'Unconfirmed audio page must not be created'
        continue
    record=inspect(row['path'])
    if row['id'] != 'MAYBACH':
        assert record['title']==row['title'], f'{row["id"]}: title mismatch'
        assert record['h1']==row['h1'], f'{row["id"]}: H1 mismatch'
    record['briefId']=row['id']
    records.append(record)
extra=['/ar/tuning','/ar/blog/mercedes-benz-maintenance-guide-dubai','/blog/mercedes-amg-gt-tuning-dubai']
for path in extra: records.append(inspect(path))
for slug in ['g63','c-class','e-class','s-class']:
    path=f'/ar/blog/mercedes-{slug}-service-dubai-guide'
    record=inspect(path,False)
    assert not record['alternates'], 'Untranslated template must not advertise alternate language'
    records.append(record)

titles=[r['title'] for r in records]; descriptions=[r['description'] for r in records]
assert len(titles)==len(set(titles)), 'Duplicate affected-page titles'
assert len(descriptions)==len(set(descriptions)), 'Duplicate affected-page descriptions'

with (OUT/'query-coverage.csv').open('w',newline='',encoding='utf-8-sig') as f:
    writer=csv.DictWriter(f,fieldnames=list(PLAN['queries'][0]))
    writer.writeheader();writer.writerows(PLAN['queries'])
result={'passed':True,'briefs':31,'queries':353,'originalMetricsUnchanged':True,
    'workbookTabsRead':{name:len(rows)-1 for name,rows in tabs.items()},
    'pagesChecked':len(records),'records':records,
    'notes':['Maybach retains its existing service/repair title and H1; its change is contextual navigation.',
             'FAQ questions are checked in initial HTML; collapsed answers use the same content objects and representative browser expansion is checked separately.',
             'No current ranking or query-to-ranking-page claim is inferred from the historical export.']}
(OUT/'generated-html-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k not in ['records','notes']},ensure_ascii=False))
