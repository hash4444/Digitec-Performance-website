"""Read-only reopening of original GSC workbooks. Never saves a source workbook."""
from pathlib import Path
import hashlib, json, re, warnings, datetime
import openpyxl

warnings.filterwarnings('ignore', message='Workbook contains no default style')
OUT = Path(__file__).resolve().parent
DOWNLOADS = Path(r'C:\Users\ADMIN\Downloads')
QUERY = 'mercedes service dubai'
TARGET = 'https://digitecme.com/brands/mercedes-benz-service-dubai'

def norm(v): return re.sub(r'\s+', ' ', str(v).strip().lower())
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def json_default(v):
    if isinstance(v, (datetime.date, datetime.datetime)): return v.isoformat()
    raise TypeError(type(v).__name__)
def write(name,v): (OUT/name).write_text(json.dumps(v,ensure_ascii=False,indent=2,default=json_default),encoding='utf-8')

manifest=[]; exact=[]; supporting=[]; independent=[]; true_filtered=[]; contains_filtered=[]
data_signatures={}; checks=0
for p in sorted(DOWNLOADS.glob('digitecme.com-Performance-on-Search-2026-*.xlsx')):
    before=sha(p)
    wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
    data={}; ranges={}
    for s in wb:
        rows=list(s.values); headers=rows[0] if rows else []
        data[s.title]=[(i,dict(zip(headers,r))) for i,r in enumerate(rows[1:],2) if any(v is not None for v in r)]
        ranges[s.title]=f'A1:{openpyxl.utils.get_column_letter(s.max_column)}{s.max_row}'
    wb.close()
    assert sha(p)==before; checks+=1
    filters={r.get('Filter'):r.get('Value') for _,r in data.get('Filters',[])}
    time_key=next((k for _,r in data.get('Chart',[]) for k in r if k=='Date' or k.startswith('Time (')),None)
    dates=[str(r[time_key]) for _,r in data.get('Chart',[]) if time_key and r.get(time_key)]
    serial={k:[v for _,v in vals] for k,vals in data.items()}
    table_sha=hashlib.sha256(json.dumps(serial,ensure_ascii=False,sort_keys=True,default=json_default).encode()).hexdigest()
    src={'name':p.name,'path':str(p),'sha256':before,'source_unchanged':True,'filters':filters,'start':min(dates) if dates else None,'end':max(dates) if dates else None,'time_column':time_key,'time_rows':len(dates),'sheets':{k:{'range':ranges[k],'rows':len(v)} for k,v in data.items()},'semantic_data_sha256':table_sha,'semantic_duplicate_of':data_signatures.get(table_sha)}
    data_signatures.setdefault(table_sha,p.name)
    manifest.append(src)
    base={'source_name':p.name,'source_path':str(p),'filters':filters,'start':src['start'],'end':src['end'],'time_column':time_key,'source_sha256':before,'semantic_duplicate_of':src['semantic_duplicate_of']}
    for i,r in data.get('Queries',[]):
        q=r.get('Top queries','')
        evidence={**base,'sheet':'Queries','range':f'A{i}:E{i}',**r,'attribution':'Exact query aggregate across all pages/countries/devices except explicit source filters. Separate workbook tabs cannot be joined.'}
        if norm(q)==QUERY: exact.append(evidence)
        if ('mercedes' in norm(q) or 'مرسيدس' in norm(q)) and re.search(r'servic|maintenance|oil|assyst|صيانة|زيت',norm(q)):
            supporting.append(evidence)
    query_filter=str(filters.get('Query',''))
    if norm(query_filter)==QUERY:
        true_filtered.append({'source':src,'exact_query_filtered_dimensions':{s:[{'row':i,**r} for i,r in data.get(s,[])] for s in ['Queries','Pages','Countries','Devices','Chart']},'filter_semantics':'Query label alone does not state equality versus contains. Must inspect Queries tab and original filter settings to confirm exact-match operator.'})
    if norm(query_filter)=='+'+QUERY:
        contains_filtered.append({'source':src,'contains_query_filtered_dimensions':{s:[{'row':i,**r} for i,r in data.get(s,[])] for s in ['Queries','Pages','Countries','Devices','Chart']},'filter_semantics':'Plus prefix denotes contains filter. A sole visible exact query row does not prove every omitted/anonymized query matches exactly. Live exact-query report independently observed same headline/page/country totals.'})
    for s in ['Pages','Countries','Devices']:
        for i,r in data.get(s,[]):
            keep=s=='Devices' or (s=='Countries' and norm(r.get('Country','')) in ['united arab emirates','are','uae']) or (s=='Pages' and ('mercedes' in str(r.get('Top pages','')).lower() or str(r.get('Top pages',''))==TARGET))
            if keep:
                independent.append({**base,'sheet':s,'range':f'A{i}:E{i}',**r,'attribution':'Independent sheet aggregate. Does not segment exact target query unless source explicitly filters that exact query. Do not multiply, allocate or join with Queries rows.'})

primary=next(r for r in exact if r['source_name']=='digitecme.com-Performance-on-Search-2026-10-12 (1).xlsx')
assert (primary['Clicks'],primary['Impressions'],primary['Position'])==(2,613,13.88)
assert (primary['start'],primary['end'])==('2026-07-10','2026-10-09')
checks+=2
primary['calculated_ctr']=primary['Clicks']/primary['Impressions']
assert round(primary['calculated_ctr']*100,2)==0.33;checks+=1
latest=max(manifest,key=lambda r:r['name'].split('2026-')[-1])
latest_exact=[r for r in exact if r['source_name']==latest['name']]
assert len(manifest)>=17;checks+=1

write('source-manifest.json',manifest)
write('exact-query-observations.json',exact)
write('mercedes-servicing-supporting-query-observations.json',supporting)
write('independent-page-country-device-evidence.json',independent)
write('exact-query-filtered-evidence.json',true_filtered)
write('contains-query-filtered-evidence.json',contains_filtered)
summary={
    'as_of':'2026-10-20 Asia/Dubai','exact_query':QUERY,'editorial_target_url':TARGET,
    'primary_baseline':primary,'latest_available_export':latest,'latest_available_exact_query_rows':latest_exact,
    'raw_workbook_count':len(manifest),'exact_query_source_observations_not_summed':len(exact),
    'exact_query_filtered_workbooks':len(true_filtered),'original_workbooks_reopened_read_only':True,
    'all_source_hashes_unchanged':all(r['source_unchanged'] for r in manifest),'checks':checks,
    'evidence_limits':{
        'exact_query_by_page':'No exact target-query-filtered Pages export and no joined query/page API response supplied.' if not true_filtered else 'Historical exact-query-labelled workbook available. See exact-query-filtered-evidence.json and its filter-operator caveat. Latest live evidence is integrated below when available.',
        'exact_query_uae':'Unavailable. UAE country totals in an unfiltered/contains-mercedes workbook are not the exact target query.',
        'exact_query_mobile_desktop':'Unavailable. Device totals in an unfiltered/contains-mercedes workbook are not the exact target query.',
        'google_selected_canonical':'Unavailable. Requires authenticated URL Inspection indexed report; HTML self-canonical alone cannot establish it.',
        'actual_query_landing_page':'Unverified for exact target query. Repository routing and self-canonical establish intended owner, not Google attribution.',
        'latest_range':f"Latest available workbook {latest['name']}, period {latest['start']} to {latest['end']}. Filename is not reporting period.",
        'source_missing_rows':'Absent from visible exported query rows means unavailable, never zero demand.',
        'overlap':'Date windows overlap and duplicate exports exist. No observation totals pooled and no causal trend inferred.',
        'position':'GSC average position 13.88 is an aggregate historical metric; not current Dubai rank, rank on every impression or evidence of top three.',
        'conversion':'GSC clicks do not establish qualified enquiries or bookings.'
    },
    'access_check':{'enabled_gsc_or_url_inspection_tools':False,'initial_browser_inventory':'Only in-app browser and MCP Apps surfaces, initially both with zero tabs. Live owned-site GSC card later found by SERP agent, read-only follow-up pending.'}
}
live_path=OUT/'live-gsc-evidence.json'
if live_path.exists():
    live=json.loads(live_path.read_text(encoding='utf-8'))
    summary['live_evidence']=live
    summary['evidence_limits'].update({
        'exact_query_by_page':'VERIFIED LIVE. Exact query 20 September–17 October 2026, Pages tab and UAE+exact hub filters inspected directly. See live-gsc-evidence.json.',
        'exact_query_uae':'VERIFIED LIVE. Exact query + United Arab Emirates. 179 impressions, 2 clicks, average position 12.2 shown. Do not confuse with all-country baseline.',
        'exact_query_mobile_desktop':'VERIFIED LIVE. Query + UAE + exact target hub: Desktop 100 impressions/2 clicks/11.2 position; Mobile 61/0/11.3. Device and page metrics jointly filtered, not allocated from independent tabs.',
        'google_selected_canonical':'VERIFIED indexed URL Inspection. Google-selected canonical = Inspected URL, exact target hub. Page indexed, fetch/crawl/indexing allowed. Last crawl displayed 16 Oct 2026 21:18:38, timezone not exposed.',
        'actual_query_landing_page':'VERIFIED LIVE. Intended hub 164 impressions/2 clicks globally; 161/2 for UAE, current 28-day exact query. Dominant observed owner, not exclusive across all time.',
        'release_effect':'Unknown. Current 28-day report ends 17 October; indexed crawl dated 16 October. Latest live provisional 24-hour span 19–20 October contains only four impressions. No release causation inferred.'
    })
    summary['access_check']['live_authenticated_gsc_access']=True
    summary['access_check']['account_writes']=False
write('exact-query-diagnosis.json',summary)
print(json.dumps({'raw_workbooks':len(manifest),'assertions':checks,'exact_rows':len(exact),'exact_query_filtered_books':len(true_filtered),'exact_observations':[{'source':r['source_name'],'period':[r['start'],r['end']],'clicks':r['Clicks'],'impressions':r['Impressions'],'ctr':r['CTR'],'position':r['Position'],'range':r['range'],'filters':r['filters']} for r in exact],'latest_exact':latest_exact},indent=2,ensure_ascii=False))
