"""Read-only repair demand extraction. Source workbooks are never saved or modified."""
from pathlib import Path
import csv
import hashlib
import json
import re
import unicodedata
import warnings

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
SOURCES = Path(r'C:\Users\ADMIN\Downloads')
PRIOR = ROOT / 'outputs/five-brand-service-2026-10-19/gsc'
warnings.filterwarnings('ignore', message='Workbook contains no default style')

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def norm(value):
    return re.sub(r'\s+', ' ', re.sub(r'[-\u2013\u2014]', ' ', unicodedata.normalize('NFKC', value).lower())).strip()

def metric_total(rows):
    clicks = sum(float(r.get('Clicks') or 0) for r in rows)
    impressions = sum(float(r.get('Impressions') or 0) for r in rows)
    return {
        'clicks': clicks, 'impressions': impressions,
        'ctr': clicks / impressions if impressions else None,
        'impression_weighted_position': sum(float(r.get('Position') or 0) * float(r.get('Impressions') or 0) for r in rows) / impressions if impressions else None,
    }

prior = {Path(r['path']).name: r for r in json.loads((PRIOR / 'source-manifest.json').read_text(encoding='utf-8'))}
manifest, books = [], {}
for path in sorted(SOURCES.glob('digitecme.com-Performance-on-Search-2026-*.xlsx')):
    before = sha(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    data, ranges = {}, {}
    for sheet in wb:
        values = list(sheet.values)
        headers = list(values[0]) if values else []
        data[sheet.title] = [dict(zip(headers, row)) for row in values[1:] if any(v is not None for v in row)]
        ranges[sheet.title] = f'A1:{openpyxl.utils.get_column_letter(sheet.max_column)}{sheet.max_row}'
    wb.close()
    after = sha(path)
    assert before == after, f'Source changed: {path}'
    filters = {r.get('Filter'): r.get('Value') for r in data.get('Filters', [])}
    time_key = next((key for key in (data.get('Chart') or [{}])[0] if key == 'Date' or key.startswith('Time (')), None)
    dates = [str(r[time_key]) for r in data.get('Chart', []) if time_key and r.get(time_key)]
    source = {
        'path': str(path), 'sha256_before': before, 'sha256_after': after,
        'source_unchanged': before == after,
        'matches_prior_sprint_hash': before == prior.get(path.name, {}).get('sha256'),
        'filters': filters, 'time_column': time_key,
        'start': min(dates) if dates else None, 'end': max(dates) if dates else None,
        'time_rows': len(dates),
        'sheets': {name: {'range': ranges[name], 'rows': len(rows), **(metric_total(rows) if name != 'Filters' else {})} for name, rows in data.items()},
        'extraction': 'openpyxl read_only=True/data_only=True; no workbook save',
    }
    table_hash = hashlib.sha256(json.dumps(data, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    source['data_table_sha256'] = table_hash
    manifest.append(source)
    books[path.name] = {'source': source, 'data': data}
assert len(manifest) == 17, f'Expected 17 source workbooks; found {len(manifest)}'
assert all(r['matches_prior_sprint_hash'] for r in manifest), 'Review updated source before reusing prior baseline'

PRIMARY = {
    '24h': 'digitecme.com-Performance-on-Search-2026-10-19.xlsx',
    '28d_historical': 'digitecme.com-Performance-on-Search-2026-09-05.xlsx',
    '92d': 'digitecme.com-Performance-on-Search-2026-10-12 (1).xlsx',
}

brand_patterns = {
    'Mercedes-Benz': r'\bmercedes\b|\bamg\b|مرسيدس',
    'BMW': r'\bbmw\b|بي ام دبليو|بي إم دبليو',
    'Rolls-Royce': r'\brolls\s*royce\b|رولز|رولس',
    'Lamborghini': r'\blamborghini\b|لامبور',
    'Volkswagen': r'\bvolkswagen\b|\bvw\b|فولكس',
}
slugs = {'Mercedes-Benz': 'mercedes-benz', 'BMW': 'bmw', 'Rolls-Royce': 'rolls-royce', 'Lamborghini': 'lamborghini', 'Volkswagen': 'volkswagen'}

def brand_for(query):
    return next((brand for brand, pattern in brand_patterns.items() if re.search(pattern, norm(query))), None)

def intent_for(query):
    q = norm(query)
    if sum(bool(re.search(pattern, q)) for pattern in brand_patterns.values()) > 1:
        return 'Ambiguous mixed-brand query; retain evidence, exclude from commercial targeting'
    if 'abu dhabi' in q:
        return 'Other-city query; exclude from Dubai priority ranking'
    if re.search(r'device detected|too hot|gearbox problems', q):
        return 'Informational warning/problem query; possible supporting answer, not direct service-intent proof'
    return 'Commercial repair/diagnostic service intent'

def category_for(query):
    q = norm(query)
    if re.search(r'tun(?:ing|ning)|remap|\bstage\b|chip|performance exhaust|exhaust upgrade|vin decoder|for sale|forsale|\bjobs\b|\bsalary\b|\btraining\b|\bcourse\b', q):
        return None
    if re.search(r'\b(?:diagnos\w*|xentry|star diagnosis)\b|فحص الكمبيوتر|تشخيص', q):
        return 'diagnostics'
    if re.search(r'\b(?:transmission|gearbox)\b|ناقل|جير', q) and re.search(r'repair|service|problem|تصليح|إصلاح|صيانة', q):
        return 'transmission-repair'
    if re.search(r'\b(?:suspension|airmatic)\b|تعليق|هيدروليك', q) and re.search(r'repair|service|problem|تصليح|إصلاح|صيانة', q):
        return 'suspension-repair'
    if re.search(r'\b(?:brake|brakes)\b|فرامل', q) and not re.search(r'shooting brake', q) and re.search(r'repair|service|problem|تصليح|إصلاح|صيانة', q):
        return 'brake-service'
    if re.search(r'\bac\b|air conditioning|air conditioner|مكيف|تكييف', q) and re.search(r'repair|service|تصليح|إصلاح|صيانة', q):
        return 'ac-service'
    if re.search(r'\b(?:electrical|electric)\b|كهربا', q) and re.search(r'repair|problem|تصليح|إصلاح', q):
        return 'electrical-repair'
    if re.search(r'\b(?:engine|coolant|overheat\w*|oil leak|water pump|radiator|cooling)\b|محرك', q) and re.search(r'repair|service|problem|too hot|leak|تصليح|إصلاح|صيانة', q):
        return 'mechanical-repair'
    return None

owner_lookup = {}
for rel, field in [
    ('outputs/b1/keyword-owner-records.csv', 'keyword_raw'),
    ('outputs/b3/bmw-keyword-owner-records.csv', 'raw_keyword'),
    ('outputs/b4/lamborghini-keyword-owner-records.csv', 'raw_keyword'),
    ('outputs/b5/rolls-royce-keyword-owner-records.csv', 'raw_keyword'),
    ('outputs/b8/volkswagen-keyword-owner-records.csv', 'raw_keyword'),
]:
    with (ROOT / rel).open(encoding='utf-8-sig', newline='') as stream:
        for row in csv.DictReader(stream):
            if row.get(field) and row.get('primary_owner'):
                owner_lookup.setdefault(norm(row[field]), set()).add(row['primary_owner'])

def owner_for(query):
    if intent_for(query).startswith('Ambiguous'):
        return '', 'Mixed-brand source query requires manual intent review; no commercial owner assigned'
    choices = sorted(owner_lookup.get(norm(query), []))
    if len(choices) == 1:
        return choices[0], 'Existing editorial ownership register; not observed query-to-ranking-page attribution'
    brand, category = brand_for(query), category_for(query)
    if not brand or not category:
        return '', 'No owner assigned'
    if brand == 'Mercedes-Benz':
        category_slug = {'ac-service': 'ac-repair', 'brake-service': 'brake-repair', 'diagnostics': 'diagnostics'}.get(category, category)
        owner = f'/services/mercedes-{category_slug}-dubai'
    else:
        category_slug = {'ac-service': 'ac-repair', 'brake-service': 'brake-repair', 'diagnostics': 'engine-diagnostics'}.get(category, category)
        owner = f'/brands/{slugs[brand]}-service-dubai/{category_slug}'
    if re.search(r'[\u0600-\u06ff]', query):
        owner = '/ar' + owner
    return owner, 'Editorial candidate from existing route family; parent route/capability audit must verify'

all_rows, relevant_queries = [], set()
for name, book in books.items():
    source = book['source']
    for rownum, row in enumerate(book['data'].get('Queries', []), 2):
        query = row['Top queries']
        brand, category = brand_for(query), category_for(query)
        if not brand or not category:
            continue
        relevant_queries.add(query)
        owner, basis = owner_for(query)
        all_rows.append({
            'query': query, 'brand': brand, 'repair_category': category, 'intent_review': intent_for(query),
            'source_file': name, 'source_path': source['path'],
            'source_range': f'Queries!A{rownum}:E{rownum}',
            'date_start': source['start'], 'date_end': source['end'],
            'filters': json.dumps(source['filters'], ensure_ascii=False),
            'period_label': next((label for label, primary in PRIMARY.items() if primary == name), 'other_historical_or_filtered'),
            'clicks': row['Clicks'], 'impressions': row['Impressions'], 'ctr': row['CTR'], 'position': row['Position'],
            'evidence_class': 'Measured original GSC query',
            'editorial_owner': owner, 'owner_basis': basis,
            'observed_ranking_url': 'Unavailable from separate Queries and Pages tabs',
            'data_table_sha256': source['data_table_sha256'],
        })

requested = [
    'mercedes engine repair dubai', 'mercedes suspension repair dubai', 'mercedes diagnostics dubai',
    'bmw transmission repair dubai', 'bmw engine repair dubai',
    'rolls royce engine repair dubai', 'rolls royce suspension repair dubai',
    'lamborghini transmission repair dubai', 'lamborghini engine diagnostics dubai',
]
proposed = set(requested) - {norm(q) for q in relevant_queries}
queries = sorted(relevant_queries | set(requested))
query_indexes = {
    label: {norm(row['Top queries']): (n, row) for n, row in enumerate(books[name]['data']['Queries'], 2)}
    for label, name in PRIMARY.items()
}

def period_metrics(label, query):
    match = query_indexes[label].get(norm(query))
    return {
        label + '_observed': bool(match),
        label + '_clicks': match[1]['Clicks'] if match else None,
        label + '_impressions': match[1]['Impressions'] if match else None,
        label + '_ctr': match[1]['CTR'] if match else None,
        label + '_position': match[1]['Position'] if match else None,
        label + '_source_file': PRIMARY[label],
        label + '_range': f'Queries!A{match[0]}:E{match[0]}' if match else None,
        label + '_missing_meaning': '' if match else 'Not present in visible exported query rows; not measured zero demand',
    }

matrix = []
for query in queries:
    owner, basis = owner_for(query)
    row = {
        'query': query, 'brand': brand_for(query), 'repair_category': category_for(query), 'intent_review': intent_for(query),
        'requested_query': norm(query) in requested,
        'evidence_class': 'Proposed keyword variant; no exact query row in the 17 reviewed exports' if norm(query) in proposed else 'Measured original GSC query in at least one reviewed export',
        'editorial_owner': owner, 'owner_basis': basis,
        'observed_ranking_url': 'Unavailable; no query x page joined export supplied',
        'competition': 'Not measured by GSC; parent qualitative live SERP audit required',
        'commercial_relevance': 'Repair or diagnostic service intent; explicit Dubai variant' if 'dubai' in norm(query) else 'Repair/diagnostic intent; location and commercial specificity need query review',
    }
    for label in PRIMARY:
        row.update(period_metrics(label, query))
    row['historical_evidence_count_no_summing'] = len([r for r in all_rows if norm(r['query']) == norm(query)])
    matrix.append(row)

top_queries = [
    ('mercedes engine repair dubai', 'P1', 'Highest measured exact local repair demand (271 impressions), page-two boundary (10.78), existing mechanical owner.'),
    ('mercedes suspension repair dubai', 'P1', '235 impressions at 13.60 with no clicks; existing focused suspension owner and repeated historical demand.'),
    ('mercedes transmission repair dubai', 'P1', 'Discovered query: 131 impressions, one click, position 15.23; proven commercial click signal and focused transmission owner.'),
    ('bmw transmission repair dubai', 'P1', '90 impressions at 11.60; localized repair intent with existing transmission owner; historical 28-day row corroborates.'),
    ('bmw engine repair dubai', 'P1', '74 impressions at 16.35; existing mechanical owner; 24-hour row has four impressions at position 6.00, too small to infer sustained improvement.'),
    ('mercedes diagnostic dubai', 'P1', 'Measured singular wording: 63 impressions at 6.60 and zero clicks; snippet/CTA relevance opportunity on established diagnostics owner.'),
    ('rolls royce engine repair dubai', 'P2', '115 impressions at 26.01; valuable explicit local repair demand but larger ranking gap than P1 opportunities.'),
    ('lamborghini engine repair dubai', 'P2', 'Discovered query: 70 impressions at 25.16; supported mechanical owner and five fresh 24-hour impressions at 11.20.'),
    ('rolls royce suspension repair dubai', 'P2', '66 impressions at 19.33; focused existing suspension owner, distinct repair intent.'),
    ('lamborghini transmission repair dubai', 'P2', '49 impressions at 15.14; existing focused owner and current 24-hour impression; clear local commercial service intent.'),
]

# Competition is primary-provider page evidence, not a localized Google rank or KD estimate.
competition = {
    '/services/mercedes-mechanical-repair-dubai': (
        'https://carzillauae.com/mercedes-engine-repair-or-rebuild/',
        'https://bavarland.ae/engine-repair.html',
        'Dedicated engine-repair providers cover this commercial task. Prioritize clear symptom-to-inspection explanations and repair options on the existing mechanical owner; no separate engine keyword route is justified.'),
    '/services/mercedes-suspension-repair-dubai': (
        'https://alkasergarage.com/mercedes-suspension-repair/',
        'https://www.avenueautoservices.com/brands/mercedes/mercedes-benz-suspension-repair-dubai/',
        'Focused Mercedes suspension pages compete for the same task. Make suspension symptoms, fitted-system limits, assessment and quotation steps easy to evaluate on the established owner.'),
    '/services/mercedes-transmission-repair-dubai': (
        'https://mercedesrepairsdubai.ae/mercedes-gearbox-repair-dubai/',
        'https://www.ityrecare.com/mercedes-transmission-repair-dubai',
        'Dedicated gearbox/transmission pages establish distinct commercial competition. Explain the inspection-to-quote sequence and confirmed gearbox scope while grouping gearbox and transmission wording on one owner.'),
    '/brands/bmw-service-dubai/transmission-repair': (
        'https://carsnowautoservices.ae/bmw-repair-experts/bmw-transmission-repair',
        'https://www.powerworksgarage.com/car-servicing-dubai/transmission-repair-dubai/bmw',
        'BMW transmission pages provide service breakdowns and direct enquiry paths. Existing near-page-one query visibility supports clearer diagnosis-versus-fluid-service explanations and vehicle-specific quotation intake on the current owner.'),
    '/brands/bmw-service-dubai/mechanical-repair': (
        'https://autorepairuae.com/bmw-repair-dubai/engine-repair/',
        'https://kaimanauto.com/engine-repair/bmw-engine-repair/',
        'Dedicated BMW engine-repair pages exist. Strengthen verified engine symptom assessment and quote clarity on mechanical repair; avoid unsupported rebuild, compatibility or turnaround claims.'),
    '/services/mercedes-diagnostics-dubai': (
        'https://mbcod.com/mercedes-diagnostic-dubai',
        'https://mercedesrepairsdubai.ae/mercedes-computer-diagnostics-ecu-coding-dubai/',
        'Competing diagnostic pages describe scan interpretation and enquiry steps. Differentiate through a precise explanation that codes guide further testing, supported functions depend on the vehicle, and repair approval follows findings.'),
    '/brands/rolls-royce-service-dubai/mechanical-repair': (
        'https://orangeauto.ae/brands/rolls-royce/engine-repair-service/',
        'https://munichmotorworks.ae/rolls-royce-repair/rolls-royce-engine-repair/',
        'Brand-specific engine-repair competition exists. With average position 26.01, preserve a longer-term priority: deepen verified assessment and quotation content on the current mechanical owner; no short-term ranking gain is forecast.'),
    '/brands/rolls-royce-service-dubai/suspension-repair': (
        'https://rollsroycerepair.ae/repair-services/rolls-royce-suspension-repair-shop-in-dubai/',
        'https://www.avenueautoservices.com/brands/rolls-royce/rolls-royce-suspension-repair-dubai/',
        'Dedicated Rolls-Royce suspension pages exist. Clarify ride-height, noise and leak assessment only within documented capability, then make vehicle details and quote requests easy to submit.'),
    '/brands/lamborghini-service-dubai/mechanical-repair': (
        'https://munichmotorworks.ae/lamborghini-repair/lamborghini-engine-repair/',
        'https://quickfitautos.com/car-brands/lamborghini/lamborghini-engine-repair-dubai/',
        'Dedicated Lamborghini engine-repair pages compete for the task. Existing mechanical repair should retain ownership; improve symptom assessment and confirmed repair scope without inferring engine-diagnostic query demand.'),
    '/brands/lamborghini-service-dubai/transmission-repair': (
        'https://quickfitautos.com/car-brands/lamborghini/lamborghini-transmission-repair/',
        'https://munichmotorworks.ae/lamborghini-repair/lamborghini-transmission-repair/',
        'Lamborghini transmission-specific commercial pages exist. Retain the current owner and explain vehicle/gearbox identification, inspection, findings and quotation, with compatibility confirmed case by case.'),
}

owner_audit_path = ROOT / 'outputs/repair-service-2026-10-19/audit/baseline-repair-owner-audit.json'
owner_audit = json.loads(owner_audit_path.read_text(encoding='utf-8'))
audit_by_route = {record['route']: record for record in owner_audit['records']}
for row in matrix:
    owner = row['editorial_owner']
    en_owner = owner[3:] if owner.startswith('/ar/') else owner
    en = audit_by_route.get(en_owner)
    ar = audit_by_route.get('/ar' + en_owner)
    row.update({
        'owner_audit_source': str(owner_audit_path) if en else None,
        'owner_audit_basis': owner_audit['basis'] if en else 'Historical editorial register or route-family candidate; outside the focused owner audit',
        'en_owner_manifest_indexable': en['manifestIndexable'] if en else None,
        'ar_owner_manifest_indexable': ar['manifestIndexable'] if ar else None,
        'en_owner_self_canonical': en['selfCanonical'] if en else None,
        'ar_owner_self_canonical': ar['selfCanonical'] if ar else None,
        'en_owner_in_sitemap': en['inSitemap'] if en else None,
        'ar_owner_in_sitemap': ar['inSitemap'] if ar else None,
        'en_hub_owner_inlinks': len(en['hubOwnerLinks']) if en else None,
        'ar_hub_owner_inlinks': len(ar['hubOwnerLinks']) if ar else None,
        'competition': 'Dedicated brand/service competition observed on primary business pages; qualitative assessment only' if en_owner in competition else 'Not assessed for this query/service in focused competition review',
        'competitor_url_1': competition[en_owner][0] if en_owner in competition else None,
        'competitor_url_2': competition[en_owner][1] if en_owner in competition else None,
        'competition_implication': competition[en_owner][2] if en_owner in competition else None,
        'competition_observed_date': '2026-10-19' if en_owner in competition else None,
        'competition_method': 'Primary provider pages reviewed by parent; supplemental second BMW transmission and Mercedes diagnostic sources opened directly. No Google Dubai rank, market volume or keyword difficulty score measured.' if en_owner in competition else None,
        'commercial_rank': None,
        'priority': None,
        'ranking_status': 'Unranked supporting keyword evidence; do not treat source impressions as market volume',
    })
    if en:
        row['owner_basis'] = 'Existing commercial owner verified in local route manifest, rendered HTML, canonical, sitemap, English/Arabic counterpart and brand-hub inlinks; not observed GSC query-to-ranking-URL attribution'

top10 = []
for rank, (query, priority, reason) in enumerate(top_queries, 1):
    item = dict(next(row for row in matrix if norm(row['query']) == query))
    item.update({'commercial_rank': rank, 'priority': priority, 'ranking_reason': reason,
                 'ranking_status': 'Final editorial implementation priority combining observed GSC demand/position, verified existing ownership, commercial relevance and qualitative provider-page competition; not a traffic, ranking or enquiry forecast'})
    top10.append(item)
    source_row = next(row for row in matrix if norm(row['query']) == query)
    source_row.update({'commercial_rank': rank, 'priority': priority, 'ranking_status': item['ranking_status']})

checked_cells = 0
for row in matrix:
    for label in PRIMARY:
        match = query_indexes[label].get(norm(row['query']))
        if match:
            assert all(row[label + '_' + key] == match[1][field] for key, field in [('clicks', 'Clicks'), ('impressions', 'Impressions'), ('ctr', 'CTR'), ('position', 'Position')])
            checked_cells += 4
        else:
            assert all(row[label + '_' + key] is None for key in ['clicks', 'impressions', 'ctr', 'position', 'range'])
assert len([r for r in matrix if r['requested_query']]) == 9
assert all(r['92d_observed'] for r in top10)
assert all(r['intent_review'] == 'Commercial repair/diagnostic service intent' for r in top10)
assert all(r['en_owner_manifest_indexable'] and r['ar_owner_manifest_indexable'] and r['en_owner_self_canonical'] and r['ar_owner_self_canonical'] and r['en_owner_in_sitemap'] and r['ar_owner_in_sitemap'] and r['en_hub_owner_inlinks'] > 0 and r['ar_hub_owner_inlinks'] > 0 for r in top10)
assert all(r['competitor_url_1'] and r['competitor_url_2'] and r['competition_implication'] for r in top10)
verification = {
    'source_workbooks_fresh_read': len(manifest),
    'source_sha256_pre_post_and_prior_sprint_match': all(r['source_unchanged'] and r['matches_prior_sprint_hash'] for r in manifest),
    'primary_period_source_rows_metric_cells_reconciled': checked_cells,
    'unobserved_query_metrics_preserved_null': True,
    'requested_query_rows': 9,
    'top10_all_have_measured_92day_query_rows': True,
    'top10_excludes_flagged_informational_mixed_brand_other_city_queries': True,
    'top10_existing_en_ar_owners_indexable_self_canonical_in_sitemap_with_hub_inlinks': True,
    'top10_two_primary_provider_competition_urls_and_qualitative_implications': True,
    'competition_not_google_local_rank_or_keyword_difficulty': True,
    'separate_queries_pages_tabs_never_joined': True,
    'overlapping_exports_never_summed': True,
    'read_only_extraction_no_workbook_save': True,
}

owners = {row['editorial_owner'] for row in matrix if row['editorial_owner']}
page_rows = []
for label, name in PRIMARY.items():
    source = books[name]['source']
    for rownum, row in enumerate(books[name]['data']['Pages'], 2):
        url = row['Top pages']
        if any(owner in url for owner in owners):
            page_rows.append({'period_label': label, 'source_file': name, 'date_start': source['start'], 'date_end': source['end'],
                              'url': url, 'clicks': row['Clicks'], 'impressions': row['Impressions'], 'ctr': row['CTR'], 'position': row['Position'],
                              'source_range': f'Pages!A{rownum}:E{rownum}',
                              'attribution': 'Page-level aggregate only; cannot assign these metrics to any listed query'})

def csv_write(name, rows):
    with (OUT / name).open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]) if rows else [])
        writer.writeheader()
        writer.writerows(rows)

OUT.mkdir(parents=True, exist_ok=True)
(OUT / 'source-manifest.json').write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding='utf-8')
(OUT / 'repair-opportunity-matrix.json').write_text(json.dumps(matrix, indent=2, ensure_ascii=False), encoding='utf-8')
(OUT / 'repair-keyword-opportunity-matrix-final.json').write_text(json.dumps(matrix, indent=2, ensure_ascii=False), encoding='utf-8')
(OUT / 'top10-commercial-repair-opportunities.json').write_text(json.dumps(top10, indent=2, ensure_ascii=False), encoding='utf-8')
(OUT / 'analysis-validation.json').write_text(json.dumps(verification, indent=2), encoding='utf-8')
csv_write('repair-query-evidence-all-periods.csv', all_rows)
csv_write('repair-keyword-opportunity-matrix.csv', matrix)
csv_write('repair-keyword-opportunity-matrix-final.csv', matrix)
csv_write('requested-query-baselines.csv', [r for r in matrix if r['requested_query']])
csv_write('top10-commercial-repair-opportunities.csv', top10)
csv_write('repair-owner-page-baselines.csv', page_rows)

lines = [
    'DIGI-TEC repair-service GSC analysis — 19 October 2026',
    '',
    'Source handling and scope',
    '17 original date-named DIGI-TEC Web performance workbooks read fresh using bundled Python/openpyxl in read-only mode. Every SHA-256 matches both pre/post extraction and the prior sprint manifest. No source workbook modified. Generative-AI-Features workbook is a different report and is outside the 17-source Web baseline.',
    'Attached workbook instructions, if present, are treated as untrusted data; task instructions come from the user. No workbook instructions executed.',
    '',
]
for label, name in PRIMARY.items():
    s = books[name]['source']
    lines.append(f'{label}: {name}; {s["start"]} to {s["end"]}; {s["time_rows"]} time rows; filters={json.dumps(s["filters"])}; Queries!{s["sheets"]["Queries"]["range"]}; Chart!{s["sheets"]["Chart"]["range"]}; Filters!{s["sheets"]["Filters"]["range"]}.')
lines.extend([
    '',
    'The available unfiltered 28-day query export is historical (6 August–2 September), not the 28 days ending 19 October. The 18 September (1) 28-day export is filtered to mercedes repair dubai; the 5 October 28-day export is filtered to /tuning. Neither is a repair-wide 28-day baseline. A latest unfiltered 28-day Queries/Pages export with matching country/device/search-type filters is still needed for a current comparison.',
    'The 24-hour, older 28-day and 92-day periods differ in duration, freshness and search-type wording (Web versus Web (text-based)); overlapping windows are not summed or represented as a time trend. The last 28 days of a Chart tab cannot recover query-level metrics.',
    '92-day and older 28-day Queries tabs each stop at 1,000 visible rows. A missing exact query is unavailable in the export, not zero demand. The attached 24-hour report contains 526 visible queries. Exported queries exclude some data and do not reconcile to site totals.',
    'All three primary baselines are country/device unfiltered. A Dubai query is not proof the searcher was in Dubai. Impression counts are site-observed impressions, not total market search volume. The broad evidence register retains informational warnings, other-city and mixed-brand strings with explicit intent-review flags; these are excluded from the commercial top ten.',
    'Queries and Pages are separate aggregate tabs. No query-to-page attribution or cannibalization claim is made. Ownership comes from editorial registers and route audits. No average-position change is claimed as a result of local changes.',
    'CTR is retained as exported at row level; any grouped CTR must use total clicks/total impressions, and grouped position must be impression weighted. No source period is summed with another.',
    '',
    'Requested exact keyword variants',
])
for row in [r for r in matrix if r['requested_query']]:
    parts = [f'{row["query"]}: {row["evidence_class"]}']
    for label in PRIMARY:
        parts.append(f'{label}: ' + (f'{row[label+"_clicks"]:g} clicks / {row[label+"_impressions"]:g} impressions / position {row[label+"_position"]:.2f}; {row[label+"_range"]}' if row[label+'_observed'] else 'not present in visible query rows'))
    lines.append(' | '.join(parts))
lines.extend(['', 'Final top ten editorial priorities', 'Rank considers measured demand, attainable average-position range, verified existing English/Arabic commercial owners and hub inlinks, explicit local commercial intent and qualitative dedicated-service competition. It is an implementation priority, not a traffic, ranking or enquiry forecast. No Google Dubai localized ranks or keyword-difficulty scores were measured.'])
for row in top10:
    lines.append(f'{row["commercial_rank"]}. {row["query"]} — {row["92d_impressions"]:g} impressions / {row["92d_clicks"]:g} clicks / position {row["92d_position"]:.2f}; {row["editorial_owner"]}; {row["ranking_reason"]}; {row["92d_range"]}.')
lines.extend([
    '',
    'Additional observed commercial opportunities',
    'Volkswagen AC repair: 65 impressions, position 15.91 (92-day Queries!A310:E310); localized AC variant 19 impressions at 14.74 (A910:E910). Group under the existing AC owner and verify capability; lower locality specificity than exact Dubai repairs above.',
    'BMW diagnostics Dubai: 49 impressions, position 13.06 (A417:E417); BMW brake repair Dubai: 44 at 13.23 (A457:E457); Mercedes brake repair Dubai: 40 at 7.40 (A490:E490); Mercedes AC repair Dubai: 42 impressions/one click at 19.86 (A38:E38). Preserve these existing owners as the next repair backlog.',
    'BMW transmission service (62 impressions at 18.45, A321:E321) and repair in Dubai (33 at 8.27, A588:E588) are measured variants for the same transmission owner. Do not create one route per wording.',
    'Requested Mercedes diagnostics Dubai plural has no exact row in any of the 17 reviewed exports. The measured singular mercedes diagnostic dubai is separate evidence, not an exact metric substitution. Requested Lamborghini engine diagnostics Dubai also has no exact row; keep as a proposed variant under the existing diagnostics owner pending demand/capability evidence. Lamborghini engine repair Dubai is measured and belongs to mechanical repair.',
    '',
    f'Analysis counts: {len(all_rows)} repair-related query evidence rows across 17 files (overlapping and duplicate scopes not aggregated); {len(matrix)} distinct raw/proposed query strings; {len(page_rows)} separate matching editorial-owner page rows in the three primary baselines.',
    'Deliverables are CSV/JSON/text analysis only. No XLSX was authored or source workbook resaved.',
])
(OUT / 'GSC-ANALYSIS.txt').write_text('\n'.join(lines) + '\n', encoding='utf-8')
competition_lines = [
    'DIGI-TEC repair-service competition and ownership assessment — 19 October 2026',
    '',
    'All ten ranked commercial owners have existing English/Arabic routes verified as indexable, self-canonical, in the sitemap and linked from their brand hubs in the parent local rendered-owner audit. This establishes local ownership and integration, not proof of production indexing or query-to-ranking-URL attribution.',
    f'Local ownership evidence: {owner_audit_path}',
    '',
    'Competition evidence uses primary business service pages observed on 19 October 2026. Parent reviewed the supplied provider sources; the second BMW transmission source (Powerworks) and second Mercedes diagnostics source (Mercedes Repairs Dubai) were directly opened during finalization. Dedicated competitor service content and phone/WhatsApp/appointment paths demonstrate competing commercial destinations. This review does not measure Google Dubai result positions, search market volume, backlink strength, market share or keyword difficulty.',
    'The top-ten order remains unchanged because every candidate already has a suitable owner and dedicated-service competition is present across the set. Larger measured local demand and positions around 6–20 take precedence where appropriate. Rolls-Royce/Lamborghini engine repair remain lower priorities despite commercial value because average position is around 25–26; no near-term ranking gain is forecast.',
    'Competitor certifications, equipment reach, dealer status, prices, guarantees, warranties and turnaround claims are not verified for DIGI-TEC and must not be copied. The implementation follows DIGI-TEC existing documented scope and vehicle-specific confirmation.',
    '',
]
for row in top10:
    competition_lines.extend([
        f'{row["commercial_rank"]}. {row["query"]} ({row["priority"]})',
        f'Owner: {row["editorial_owner"]}. English/Arabic indexable, self-canonical, in sitemap; hub inlinks {row["en_hub_owner_inlinks"]}/{row["ar_hub_owner_inlinks"]}.',
        f'92-day GSC: {row["92d_impressions"]:g} impressions, {row["92d_clicks"]:g} clicks, average position {row["92d_position"]:.2f}; {row["92d_source_file"]} :: {row["92d_range"]}.',
        f'Primary competitor 1: {row["competitor_url_1"]}',
        f'Primary competitor 2: {row["competitor_url_2"]}',
        f'Qualitative implication: {row["competition_implication"]}',
        f'Editorial ranking reason: {row["ranking_reason"]}',
        '',
    ])
(OUT / 'competition-assessment.txt').write_text('\n'.join(competition_lines) + '\n', encoding='utf-8')
print(json.dumps({'source_count': len(manifest), 'source_hashes_preserved': all(r['source_unchanged'] and r['matches_prior_sprint_hash'] for r in manifest), 'query_evidence_rows': len(all_rows), 'matrix_rows': len(matrix), 'proposed_requested_variants': sorted(proposed), 'top10': [{'rank': r['commercial_rank'], 'query': r['query'], 'impressions92d': r['92d_impressions'], 'position92d': r['92d_position'], 'owner': r['editorial_owner']} for r in top10]}, indent=2))
